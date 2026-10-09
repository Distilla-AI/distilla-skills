#!/usr/bin/env python3
"""
check_env.py - preflight: can this host build and recalculate the workbook?

Usage:
    python check_env.py

Standard library only, so it runs on any host with Python. It checks the packages in requirements.txt
by doing the real work on a two-cell test workbook: write it with openpyxl, recalculate it with
scripts/recalc.py (the `formulas` package), and read the value back. Prints one JSON object:

  {"python": "3.12.8", "openpyxl": "3.1.5" | null, "formulas": "1.3.4" | null,
   "can_build": true, "can_recalc": true, "mode": "full" | "build_only" | "no_workbook", "next": "..."}

Exit code 0 when the workbook can be built (full or build_only), 1 when it cannot.
Modes: full = build + recalculate + verify; build_only = deliver the workbook marked "not
recalculated or verified here" (SKILL.md step 7); no_workbook = openpyxl is missing and cannot be installed.
"""
import importlib
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))


def version(mod):
    try:
        m = importlib.import_module(mod)
        return getattr(m, "__version__", "unknown")
    except Exception:  # noqa: BLE001 - any import failure means "not usable here"
        return None


def main():
    out = {"python": sys.version.split()[0], "openpyxl": version("openpyxl"), "formulas": version("formulas"),
           "can_build": False, "can_recalc": False}
    tmp = tempfile.mkdtemp(prefix="dcf_check_")
    path = os.path.join(tmp, "check.xlsx")
    if out["openpyxl"]:
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws["A1"] = 2
            ws["A2"] = "=A1*3"
            wb.save(path)
            out["can_build"] = True
        except Exception as e:  # noqa: BLE001
            out["build_error"] = f"{type(e).__name__}: {e}"
    if out["can_build"] and out["formulas"]:
        try:
            r = subprocess.run([sys.executable, os.path.join(HERE, "recalc.py"), path],
                               capture_output=True, text=True, timeout=180)
            from openpyxl import load_workbook
            v = load_workbook(path, data_only=True).active["A2"].value
            out["can_recalc"] = (v == 6)
            if not out["can_recalc"]:
                out["recalc_error"] = (r.stderr or r.stdout)[-300:] or f"A2 read back as {v!r}, expected 6"
        except Exception as e:  # noqa: BLE001
            out["recalc_error"] = f"{type(e).__name__}: {e}"
    if out["can_build"] and out["can_recalc"]:
        out["mode"], out["next"] = "full", "Proceed: build, recalculate and verify as in SKILL.md."
    elif out["can_build"]:
        out["mode"] = "build_only"
        out["next"] = ("Try `pip install -r requirements.txt` (skill folder) if this host allows installs, then rerun this "
                       "check. Otherwise build the workbook and deliver it marked 'not recalculated or verified here' "
                       "with no value per share (SKILL.md step 7).")
    else:
        out["mode"] = "no_workbook"
        out["next"] = ("openpyxl is missing: try `pip install -r requirements.txt` (skill folder) if this host allows "
                       "installs, then rerun this check. If it cannot be installed, tell the user the workbook cannot be "
                       "built on this host and stop before the build step.")
    print(json.dumps(out, indent=1))
    sys.exit(0 if out["can_build"] else 1)


if __name__ == "__main__":
    main()
