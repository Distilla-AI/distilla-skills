#!/usr/bin/env python3
"""
check_model.py - verify a recalculated model and print the key outputs.

Usage:
    python check_model.py model.xlsx [model_inputs.json] [--all]

--all first prints base, bull and bear side by side (each recalculated on a temporary copy with the
scenario switch set), for the checkpoint; then checks the workbook as saved.

Run AFTER recalc.py or the host's recalculation tool (openpyxl-written formulas have no values until then).
Checks are done independently in Python from the recalculated values, not by trusting the
workbook's own Checks tab:
  - every year's balance sheet balances
  - historical lines tie to Distilla (revenue, EBIT, net income, total assets)
  - forecast cash is never negative
  - UFCF recomputes from its components; EV = sum PV + PV(TV); per-share math
  - WACC > g, TV share of EV, consensus variance in anchored years
Exit code 0 = pass (warnings allowed), 1 = hard failure.
"""
import json
import os
import re
import subprocess
import sys
import tempfile
from openpyxl import load_workbook

ALL = "--all" in sys.argv
ARGS = [a for a in sys.argv[1:] if a != "--all"]
path = ARGS[0]
inputs = json.load(open(ARGS[1])) if len(ARGS) > 1 else None

if ALL:
    here = os.path.dirname(os.path.abspath(__file__))
    keys = ("Value per share - perpetuity growth method", "Value per share at the brokers discount rate",
            "Upside / (downside)", "RESULT")
    rows, tmp = {}, tempfile.mkdtemp(prefix="scen_")
    for n, name in ((1, "Base"), (2, "Bull"), (3, "Bear")):
        wbf = load_workbook(path)
        ws = wbf["Assumptions"]
        hit = next(r for r in range(1, ws.max_row + 1)
                   if str(ws.cell(row=r, column=1).value or "").startswith("Active scenario"))
        ws.cell(row=hit, column=3, value=n)
        cp = os.path.join(tmp, f"scenario_{n}.xlsx")
        wbf.save(cp)
        subprocess.run([sys.executable, os.path.join(here, "recalc.py"), cp], capture_output=True, text=True)
        out = subprocess.run([sys.executable, os.path.abspath(__file__), cp] + ARGS[1:2], capture_output=True, text=True).stdout
        for k in keys:
            m = re.search(re.escape(k) + r"[^\n]*?\s(\S+)\s*$", out, re.M)
            rows.setdefault(k, {})[name] = m.group(1) if m else "-"
    print("Scenarios (each recalculated on a copy; the saved workbook is unchanged)")
    print(f"{'':58}{'Base':>14}{'Bull':>14}{'Bear':>14}")
    for k in keys:
        print(f"{k[:58]:58}" + "".join(f"{rows[k][s]:>14}" for s in ("Base", "Bull", "Bear")))
    print()
wb = load_workbook(path, data_only=True)


def find_row(ws, label, col=1):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip() == label:
            return r
    raise KeyError(f"{ws.title}: row '{label}' not found")


def series(sheet, label):
    ws = wb[sheet]
    r = find_row(ws, label)
    out = []
    for c in range(3, ws.max_column + 1):
        out.append(ws.cell(row=r, column=c).value)
    return out


def scalar(sheet, label):
    ws = wb[sheet]
    return ws.cell(row=find_row(ws, label), column=3).value


hard, warn = [], []
hdr = [c.value for c in wb["Model"][4][2:]]
n = len([h for h in hdr if h])
hdr = hdr[:n]
nh = len([h for h in hdr if str(h).endswith("A")])

if any(v is None for v in series("Model", "Revenue")[:n]):
    print("Values are empty - run recalc.py first.")
    sys.exit(1)

# 1. balance sheet balances
ta = series("Model", "Total assets")[:n]
tle = series("Model", "Total liabilities & equity")[:n]
for h, a, b in zip(hdr, ta, tle):
    if abs(a - b) > max(1.0, abs(a) * 1e-6):
        hard.append(f"Balance sheet does not balance in {h}: {a:,.1f} vs {b:,.1f}")

# 2. historical ties to Distilla
if inputs:
    raw = inputs["historical"]["raw"]
    ties = [("Revenue", "Model", "income_statement_sales"),
            ("EBIT (operating income)", "Model", "income_statement_ebit_operating_income"),
            ("Net income to common", "Model", "income_statement_net_income"),
            ("Total assets", "Model", "balance_sheet_total_assets")]
    for lab, sh, m in ties:
        got = series(sh, lab)[:nh]
        exp = raw.get(m) or []
        for i, (g, e) in enumerate(zip(got, exp)):
            if e is not None and abs(g - e) > max(1.0, abs(e) * 1e-6):
                hard.append(f"{lab} {hdr[i]} = {g:,.1f} but Distilla reports {e:,.1f}")

# 2b. history ties line-by-line to Distilla subtotals (CF) and totals (BS) with no hidden plug
if inputs:
    raw = inputs["historical"]["raw"]
    def rv(m, i):
        v = (raw.get(m) or [None] * nh)[i]
        return v or 0.0
    cf_ties = [("Cash from operations", "cash_flow_net_operating_cash_flow"),
               ("Cash from investing", "cash_flow_net_investing_cash_flow"),
               ("Cash from financing", "cash_flow_net_financing_cash_flow"),
               ("Net change in cash", "cash_flow_net_change_in_cash")]
    for lab, m in cf_ties:
        got = series("Model", lab)[:nh]
        for i in range(nh):
            tol = max(1.0, 1e-4 * abs(series("Model", "Revenue")[i]))  # Distilla source rounding
            if raw.get(m) and raw[m][i] is not None and abs(got[i] - rv(m, i)) > tol:
                warn.append(f"History: {lab} {hdr[i]} sums to {got[i]:,.0f} vs Distilla {rv(m, i):,.0f}")
    for lab in ("  Unreconciled vs Distilla (current assets)", "  Unreconciled vs Distilla (long-term assets)",
                "  Unreconciled vs Distilla (current liabilities)", "  Unreconciled vs Distilla (long-term liabilities)",
                "  Unreconciled vs Distilla (equity / other instruments)"):
        got = series("Model", lab.strip())[:nh]
        big = [(hdr[i], v) for i, v in enumerate(got) if abs(v or 0) > max(1.0, abs(ta[i]) * 0.001)]
        if big:
            warn.append(f"History: {lab.strip()} not zero: " + ", ".join(f"{h} {v:,.0f}" for h, v in big))

# 2c. forecast shape
th = (inputs or {}).get("thresholds", {})
rev = series("Model", "Revenue")[:n]
opex = series("Model", "Operating expenses (SG&A, R&D, other)")[:n]
for h, o in zip(hdr[nh:], opex[nh:]):
    if o < 0:
        hard.append(f"Operating expenses negative in {h}")
ppe = series("Model", "Net PP&E")[:n]
r0, rN = ppe[nh - 1] / rev[nh - 1], ppe[-1] / rev[-1]
if not (th.get("ppe_drift_low", 0.5) * r0 <= rN <= th.get("ppe_drift_high", 2.0) * r0):
    warn.append(f"PP&E / revenue drifts from {r0:.1%} (last actual) to {rN:.1%} (terminal) - check capex and D&A")
cash_ = series("Model", "Cash & ST investments")[:n]
c0, cN = cash_[nh - 1] / rev[nh - 1], cash_[-1] / rev[-1]
if cN > th.get("cash_build_multiple", 3.0) * c0 and cN > 0.5:
    warn.append(f"Cash builds from {c0:.0%} to {cN:.0%} of revenue - raise payout/buybacks if that is unrealistic "
                "(valuation unaffected; interest income and EPS are)")

# 3. cash never negative
cash = series("Model", "Cash & ST investments")[:n]
for h, c in zip(hdr[nh:], cash[nh:]):
    if c < -0.5:  # buybacks capped at available cash leave exactly 0 (float noise -0)
        warn.append(f"Forecast cash negative in {h} ({c:,.0f}) - reduce buybacks/dividends or add debt issuance")

# 4. DCF arithmetic
ufcf = series("DCF", "Unlevered free cash flow")[nh:n]
nopat = series("DCF", "NOPAT")[nh:n]
da = series("DCF", "(+) D&A")[nh:n]
cx = series("DCF", "(-) Capex")[nh:n]
dn = series("DCF", "(-) Increase in NWC")[nh:n]
lc = series("DCF", "(-) New leases (IFRS 16 right-of-use additions)")[nh:n]
try:  # present only when a finance arm is valued separately
    fr = series("DCF", "(+) Finance receivables growth (funded by finance-arm debt)")[nh:n]
except KeyError:
    fr = [0] * len(ufcf)
try:
    fp = series("DCF", "(+) Finance-arm leased assets growth (funded by finance-arm debt)")[nh:n]
    fr = [(a or 0) + (b or 0) for a, b in zip(fr, fp)]
except KeyError:
    pass
for i, u in enumerate(ufcf):
    if abs(u - (nopat[i] + da[i] + cx[i] + dn[i] + (lc[i] or 0) + (fr[i] or 0))) > 1:
        hard.append(f"UFCF does not recompute in {hdr[nh + i]}")
pv = series("DCF", "PV of UFCF")[nh:n]
ev = scalar("DCF", "Enterprise value")
pvtv = scalar("DCF", "PV of terminal value")
if abs(ev - (sum(pv) + pvtv)) > 1:
    hard.append("EV != sum of PV + PV of terminal value")
wacc = scalar("DCF", "WACC")
g = scalar("DCF", "Terminal growth rate")
if wacc <= g:
    hard.append(f"WACC {wacc:.2%} <= terminal growth {g:.2%}")
eqv = scalar("DCF", "Equity value")
if eqv <= 0:
    warn.append(f"Equity value is negative ({eqv:,.0f}) - a DCF is not meaningful on these assumptions; "
                "revisit margins / horizon or use another method")
tvp = scalar("DCF", "Terminal value as % of EV")
if tvp > 0.85:
    warn.append(f"Terminal value is {tvp:.0%} of EV - valuation rests mostly on the terminal assumption")
if ufcf and ufcf[-1] <= 0:
    warn.append("Terminal-year UFCF is not positive - perpetuity value is not meaningful; consider exit multiple or longer horizon")

# 5. consensus variance (base scenario only)
if scalar("Assumptions", "Active scenario (1 = Base, 2 = Bull, 3 = Bear)") == 1:
    try:
        var = series("Assumptions", "Model revenue vs consensus mean")[nh:n]
        for h, v in zip(hdr[nh:], var):
            if v not in (None, 0) and abs(v) > 0.02:
                warn.append(f"Base revenue differs from consensus mean by {v:+.1%} in {h}")
    except KeyError:
        pass

# 6. formula errors that recalc might have missed
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#"):
                hard.append(f"{ws.title}!{c.coordinate} = {c.value}")

# 7. reverse DCF must reproduce the price, and terminal growth must be backed by reinvestment
g_req = scalar("DCF", "Perpetual growth needed to justify the price (other inputs unchanged)")
ev_req = scalar("DCF", "Enterprise value implied by the share price")
sum_pv = scalar("DCF", "Sum of PV of UFCF")
tN = scalar("DCF", "Discount period for terminal value (years)")
uN = ufcf[-1]
if isinstance(g_req, (int, float)) and wacc > g_req:
    ev_back = sum_pv + uN * (1 + g_req) / (wacc - g_req) / (1 + wacc) ** tN
    if abs(ev_back - ev_req) > max(1.0, 1e-4 * abs(ev_req)):
        hard.append(f"Reverse DCF does not reproduce the price-implied EV ({ev_back:,.0f} vs {ev_req:,.0f})")
    # A result about the market, not a model problem: report as a note, never as a warning
    rf_ = (inputs or {}).get("wacc", {}).get("rf")
    if rf_ and g_req > rf_:
        print(f"Note: the share price implies perpetual growth of {g_req:.1%}, above the risk-free rate ({rf_:.1%}) - "
              "the market expects far better long-run economics than this model")
    elif g_req < 0:
        print(f"Note: the share price implies perpetual growth of {g_req:.1%} - the market assigns little value "
              "beyond the forecast period (or doubts the forecast itself)")
rr = scalar("DCF", "Reinvestment rate (net reinvestment / NOPAT)")
ronic = scalar("DCF", "Implied return on new capital (terminal growth / reinvestment rate)")
nopat_N = nopat[-1] if nopat else 0
if g > 0 and nopat_N > 0 and isinstance(rr, (int, float)):
    if rr < -0.10:
        warn.append(f"Terminal growth {g:.1%} with negative reinvestment (rate {rr:.0%}): cash flow stays well above "
                    "NOPAT forever (e.g. growing prepaid subscriptions) - is that sustainable?")
    elif rr <= 0:
        print(f"Note: terminal growth needs no net reinvestment (rate {rr:.0%}) - negative working capital; small effect")
    elif isinstance(ronic, (int, float)) and ronic < wacc:
        warn.append(f"Return on new capital {ronic:.1%} is below WACC {wacc:.1%} - terminal growth destroys value")

# ---- report
print("=" * 64)
print(f"Scenario: {wb['Assumptions']['D5'].value}")
rows_ = [("WACC", "pct"), ("Terminal growth rate", "pct"), ("Enterprise value", "num"), ("Equity value", "num"),
         ("Value per share - perpetuity growth method", "num"), ("Current share price", "num"),
         ("Consensus target price (memo)", "num"), ("Upside / (downside)", "pct"), ("Terminal value as % of EV", "pct"),
         ("Terminal EV/EBITDA implied by the perpetuity method", "x"), ("Today's market EV / FY1 EBITDA (reference only)", "x"),
         ("Perpetual growth implied if today's multiple held at the terminal year", "pct"),
         ("Perpetual growth needed to justify the price (other inputs unchanged)", "pct"),
         ("EBIT margin scale needed to justify the price (1.00 = model)", "dec"),
         ("Terminal EBIT margin needed to justify the price", "pct"),
         ("Reinvestment rate (net reinvestment / NOPAT)", "pct"),
         ("Implied return on new capital (terminal growth / reinvestment rate)", "pct")]
for lab, kind in rows_:
    v = scalar("DCF", lab)
    if not isinstance(v, (int, float)):
        out = str(v)
    else:
        out = {"pct": f"{v:.2%}", "x": f"{v:.1f}x", "dec": f"{v:.2f}", "num": f"{v:,.2f}"}[kind]
    print(f"{lab[:66]:66} {out:>16}")
try:
    fav = scalar("DCF", "Finance arm value (added in the bridge)")
    print(f"{'Finance arm value (added in the bridge)':66} {fav:>16,.2f}  at P/B "
          f"{scalar('DCF', 'Finance arm P/B used'):.2f}x")
except KeyError:
    pass
try:
    vb = scalar("DCF", "Value per share at the brokers' rate")
    bw = scalar("DCF", "as a WACC at today's weights")
    if isinstance(vb, (int, float)) and isinstance(bw, (int, float)):
        print(f"{'Value per share at the brokers discount rate (' + format(bw, '.2%') + ')':66} {vb:>16,.2f}")
except KeyError:
    pass
try:
    capped = scalar("Checks", "Years with buybacks capped by available cash (information)")
    if capped:
        print(f"Note: buybacks capped by available cash in {capped:.0f} forecast year(s)")
except KeyError:
    pass
print(f"Workbook Checks tab: {scalar('Checks', 'OVERALL')}")
print("=" * 64)
for w in warn:
    print("WARNING:", w)
for h in hard:
    print("FAIL:", h)
print("RESULT:", "FAIL" if hard else ("PASS WITH WARNINGS" if warn else "PASS"))
sys.exit(1 if hard else 0)
