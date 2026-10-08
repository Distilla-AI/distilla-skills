#!/usr/bin/env python3
"""
build_model.py - write a live-formula 3-statement model + DCF workbook.

Usage:
    python build_model.py model_inputs.json output.xlsx

model_inputs.json comes from prepare_inputs.py (after the user has confirmed / edited
the assumptions). Every forecast and valuation cell is an Excel formula; the only
hardcodes are Distilla actuals (Raw Data tab), consensus reference values, and the
blue assumption inputs.

Build order mirrors how the model works:
  Raw Data -> historical IS / BS / CF (links) -> Assumptions -> Schedules
  -> forecast IS / BS / CF (cash is the balancing item) -> DCF -> Sensitivity -> Checks -> Summary
After building, run recalc.py (or the host's spreadsheet recalculation tool), then check_model.py.
"""
import json
import sys
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

M = json.load(open(sys.argv[1]))
OUT = sys.argv[2]

HP = M["historical"]["periods"]
FP = M["forecast_periods"]
NH, NF = len(HP), len(FP)
P = HP + FP
FIRST_COL = 3  # column C
CUR = M["company"].get("reporting_currency", "")
UNITS = {"m": "mm", "b": "bn", "k": "k"}.get(M.get("units", "m"), M.get("units", "m"))
A = M["assumptions"]
RAW = M["historical"]["raw"]
FA = M.get("finance_arm") or {}
FA_ON = bool(FA.get("on"))


def fa_leased_term():
    """Leased assets held flat: subtracted from the D&A base. Built at write time, once the DCF cells exist."""
    return f"-{sref('fa_leased')}*{sref('fa_on')}" if FA_ON else ""


MH = M.get("multiple_history") or {}

# ---------------------------------------------------------------- formats & styles
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
DAYS = '0.0;(0.0);"-"'
PS = '#,##0.00;(#,##0.00);"-"'
MULT = '0.0"x";(0.0"x");"-"'
DATE = 'yyyy-mm-dd'
DEC = '0.00;(0.00);"-"'
YRS = '0.00'
FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")
GREEN = Font(name=FONT, size=10, color="008000")
BLACK = Font(name=FONT, size=10, color="000000")
BOLD = Font(name=FONT, size=10, bold=True)
HDR_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
FC_FILL = PatternFill("solid", fgColor="F2F2F2")
TOP = Border(top=Side(style="thin"))

ROWS = {}   # (sheet, key) -> row
SC = {}     # scalar key -> (sheet, "C12")
CURRENT = [None]
MODEL = "Model"
# logical statement name -> physical sheet. The three statements and schedules share one tab
# so a year is always the same column and every link is visible by scrolling down.
PHYS = {"Income Statement": MODEL, "Balance Sheet": MODEL, "Cash Flow": MODEL, "Schedules": MODEL}


def col(j):
    return get_column_letter(FIRST_COL + j)


def ref(sheet, key, j):
    r = ROWS[(sheet, key)]
    ps = PHYS.get(sheet, sheet)
    return f"{col(j)}{r}" if ps == CURRENT[0] else f"'{ps}'!{col(j)}{r}"


def rng(sheet, key, j0, j1):
    r = ROWS[(sheet, key)]
    ps = PHYS.get(sheet, sheet)
    a = f"{col(j0)}{r}:{col(j1)}{r}"
    return a if ps == CURRENT[0] else f"'{ps}'!{a}"


def sref(key):
    sheet, cell = SC[key]
    c, r = cell.rstrip("0123456789"), cell[len(cell.rstrip("0123456789")):]
    a = f"${c}${r}"
    return a if sheet == CURRENT[0] else f"'{sheet}'!{a}"


def raw(metric, j):
    return ref("Raw Data", metric, j)


def is_hist(j):
    return j < NH


def fy_label(p, j):
    return f"FY{p[:4]}{'A' if is_hist(j) else 'E'}"


# ---------------------------------------------------------------- sheet builder
class SB:
    """Collects rows first (so every row number is known), writes formulas later."""

    def __init__(self, name, title, periods=True, period_header=True):
        self.name, self.title, self.periods = name, title, periods
        self.period_header = period_header
        self.items = []
        self.groups = []
        self.r = 5

    def hdr(self, text, major=False):
        self.items.append(("major" if major else "hdr", self.r, text))
        self.r += 1

    def blank(self):
        self.r += 1

    def line(self, key, label, fmt=NUM, hist=None, fc=None, unit="", bold=False, note=None,
             top=False, key_fill=False, ns=None):
        ROWS[(ns or self.name, key)] = self.r
        self.items.append(("line", self.r, dict(key=key, label=label, fmt=fmt, hist=hist, fc=fc,
                                                unit=unit, bold=bold, note=note, top=top,
                                                key_fill=key_fill)))
        self.r += 1

    def scalar(self, key, label, value, fmt=NUM, unit="", source=None, bold=False, key_fill=False):
        SC[key] = (self.name, f"C{self.r}")
        self.items.append(("scalar", self.r, dict(key=key, label=label, value=value, fmt=fmt,
                                                  unit=unit, source=source, bold=bold,
                                                  key_fill=key_fill)))
        self.r += 1

    def text(self, label, value=None, bold=False):
        self.items.append(("text", self.r, dict(label=label, value=value, bold=bold)))
        self.r += 1

    def custom(self, fn):
        self.items.append(("custom", self.r, fn))


class Section:
    """A statement living inside a shared physical sheet: own row namespace, collapsible group."""

    def __init__(self, parent, logical, title):
        self.p, self.name = parent, logical
        if parent.items:
            parent.blank()
        parent.hdr(title, major=True)
        self.start = parent.r

    def hdr(self, text):
        self.p.hdr(text)

    def blank(self):
        self.p.blank()

    def line(self, key, *a, **k):
        self.p.line(key, *a, ns=self.name, **k)

    def close(self):
        self.p.groups.append((self.start, self.p.r - 1))


def style_for(v):
    if isinstance(v, str) and v.startswith("="):
        return GREEN if "!" in v else BLACK
    return BLUE


def write_sheet(wb, sb):
    ws = wb.create_sheet(sb.name)
    CURRENT[0] = sb.name
    ws["A1"] = sb.title
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = f"{M['company'].get('name')} ({M['company'].get('ticker')}) | {CUR} {UNITS} unless stated"
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 10
    if sb.periods:
        for j, p in enumerate(P):
            ws.column_dimensions[col(j)].width = 13
            if sb.period_header:
                c = ws.cell(row=4, column=FIRST_COL + j, value=fy_label(p, j))
                c.font, c.fill = HDR_FONT, HDR_FILL
                c.alignment = Alignment(horizontal="center")
        if sb.period_header:
            for cc in (1, 2):
                ws.cell(row=4, column=cc).fill = HDR_FILL
        if sb.period_header:
            ws.freeze_panes = ws.cell(row=5, column=FIRST_COL)
    else:
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 60
    for kind, r, it in sb.items:
        if kind == "hdr":
            c = ws.cell(row=r, column=1, value=it)
            c.font = BOLD
            last = FIRST_COL + (len(P) if sb.periods else 2)
            for cc in range(1, last):
                ws.cell(row=r, column=cc).fill = SEC_FILL
        elif kind == "major":
            c = ws.cell(row=r, column=1, value=it)
            c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
            for cc in range(1, FIRST_COL + len(P)):
                ws.cell(row=r, column=cc).fill = PatternFill("solid", fgColor="2F5597")
        elif kind == "text":
            c = ws.cell(row=r, column=1, value=it["label"])
            c.font = BOLD if it["bold"] else BLACK
            if it["value"] is not None:
                v = ws.cell(row=r, column=3, value=it["value"])
                v.font = BLACK
        elif kind == "line":
            lab = ws.cell(row=r, column=1, value=it["label"])
            lab.font = BOLD if it["bold"] else BLACK
            if it["note"]:
                lab.comment = Comment(it["note"], "model")
            ws.cell(row=r, column=2, value=it["unit"]).font = Font(name=FONT, size=9, color="7F7F7F")
            for j in range(len(P)):
                fn = it["hist"] if is_hist(j) else it["fc"]
                if fn is None:
                    continue
                v = fn(j)
                if v is None:
                    continue
                c = ws.cell(row=r, column=FIRST_COL + j, value=v)
                c.number_format = it["fmt"]
                f = style_for(v)
                c.font = Font(name=FONT, size=10, color=f.color, bold=it["bold"])
                if not is_hist(j) and f is BLUE:
                    c.fill = KEY_FILL if it["key_fill"] else FC_FILL
                if it["top"]:
                    c.border = TOP
        elif kind == "scalar":
            ws.cell(row=r, column=1, value=it["label"]).font = BOLD if it["bold"] else BLACK
            ws.cell(row=r, column=2, value=it["unit"]).font = Font(name=FONT, size=9, color="7F7F7F")
            v = it["value"]
            c = ws.cell(row=r, column=3, value=v)
            c.number_format = it["fmt"]
            f = style_for(v)
            c.font = Font(name=FONT, size=10, color=f.color, bold=it["bold"])
            if it["key_fill"]:
                c.fill = KEY_FILL
            if it["source"]:
                s = ws.cell(row=r, column=4, value=it["source"])
                s.font = Font(name=FONT, size=9, italic=True, color="595959")
        elif kind == "custom":
            it(ws)
    for a, b in sb.groups:
        if b >= a:
            ws.row_dimensions.group(a, b, outline_level=1, hidden=False)
    ws.sheet_properties.outlinePr.summaryBelow = False
    return ws


# ---------------------------------------------------------------- Raw Data
RAW_METRICS = [
    ("income_statement_sales", "Sales"),
    ("income_statement_cost_of_goods_sold_cogs_incl_d_and_a", "COGS incl. D&A"),
    ("income_statement_sg_and_a_expense", "SG&A (incl. R&D)"),
    ("income_statement_research_and_development", "R&D (memo)"),
    ("income_statement_ebit_operating_income", "EBIT"),
    ("income_statement_pretax_income", "Pretax income"),
    ("income_statement_income_taxes", "Income taxes"),
    ("income_statement_minority_interest", "Minority interest (IS)"),
    ("income_statement_net_income", "Net income"),
    ("income_statement_diluted_shares_outstanding", "Diluted shares (m)"),
    ("income_statement_dividends_per_share", "DPS"),
    ("cash_flow_depreciation_depletion_and_amortization", "D&A (cash flow)"),
    ("cash_flow_deferred_taxes", "Deferred taxes"),
    ("cash_flow_funds_from_operations", "Funds from operations"),
    ("cash_flow_changes_in_working_capital", "Change in working capital"),
    ("cash_flow_net_operating_cash_flow", "Cash from operations"),
    ("cash_flow_capital_expenditures", "Capex (total)"),
    ("cash_flow_capital_expenditures_fixed_assets", "Capex - fixed assets"),
    ("cash_flow_net_assets_from_acquisitions", "Acquisitions"),
    ("cash_flow_sale_of_fixed_assets_and_businesses", "Sale of fixed assets & businesses"),
    ("cash_flow_purchase_or_sale_of_investments", "Purchase / sale of investments"),
    ("cash_flow_other_investing_funds", "Other investing"),
    ("cash_flow_net_investing_cash_flow", "Cash from investing"),
    ("cash_flow_cash_dividends_paid", "Dividends paid"),
    ("cash_flow_repurchase_of_common_and_preferred_stock", "Share repurchases"),
    ("cash_flow_sale_of_common_and_preferred_stock", "Share issuance"),
    ("cash_flow_issuance_or_reduction_of_debt_net", "Net debt issuance"),
    ("cash_flow_repayments_of_operating_lease_liabilities", "Lease principal repayments"),
    ("cash_flow_other_financing_funds", "Other financing"),
    ("cash_flow_net_financing_cash_flow", "Cash from financing"),
    ("cash_flow_exchange_rate_effect", "FX effect"),
    ("cash_flow_net_change_in_cash", "Net change in cash"),
    ("balance_sheet_cash_and_short_term_investments", "Cash & ST investments"),
    ("balance_sheet_short_term_receivables", "Receivables"),
    ("balance_sheet_inventories", "Inventories"),
    ("balance_sheet_other_current_assets", "Other current assets"),
    ("balance_sheet_total_current_assets", "Total current assets"),
    ("balance_sheet_net_property_plant_and_equipment", "Net PP&E"),
    ("balance_sheet_total_long_term_investments", "LT investments"),
    ("balance_sheet_intangible_assets", "Intangibles incl. goodwill"),
    ("balance_sheet_deferred_tax_assets", "Deferred tax assets"),
    ("balance_sheet_other_assets", "Other long-term assets"),
    ("balance_sheet_total_assets", "Total assets"),
    ("balance_sheet_short_term_debt_and_curr_portion_long_term_debt", "ST debt"),
    ("balance_sheet_accounts_payable", "Accounts payable"),
    ("balance_sheet_income_tax_payable", "Income tax payable"),
    ("balance_sheet_other_current_liabilities", "Other current liabilities"),
    ("balance_sheet_total_current_liabilities", "Total current liabilities"),
    ("balance_sheet_long_term_debt_excl_lease_obligations", "LT debt excl. leases"),
    ("balance_sheet_capital_and_operating_lease_obligations", "Lease liabilities"),
    ("balance_sheet_provision_for_risks_and_charges", "Provisions"),
    ("balance_sheet_deferred_tax_liabilities", "Deferred tax liabilities"),
    ("balance_sheet_other_liabilities", "Other long-term liabilities"),
    ("balance_sheet_total_liabilities", "Total liabilities"),
    ("balance_sheet_total_shareholders_equity", "Shareholders' equity"),
    ("balance_sheet_accumulated_minority_interest", "Minority interest (BS)"),
]

rawsb = SB("Raw Data", "Raw Data - Distilla actuals (as delivered)")
rawsb.hdr("Source: Distilla MCP, financial_data_point, provenance = financials, annual")
for m, lab in RAW_METRICS:
    vals = RAW.get(m, [None] * NH)
    missing = [HP[i] for i, v in enumerate(vals) if v is None]
    note = f"Distilla metric: {m}" + (f"\nNot reported for {missing}; set to 0." if missing else "")
    rawsb.line(m, lab, fmt=DEC if "shares" in m or "per_share" in m or m.endswith("dividends_per_share") else NUM,
               hist=(lambda vals: lambda j: vals[j] if vals[j] is not None else 0)(vals), note=note)

# ---------------------------------------------------------------- Assumptions
asb = SB("Assumptions", "Assumptions - forecast drivers (blue = input)")
sel_row = asb.r
asb.custom(lambda ws: None)
SC["scenario"] = ("Assumptions", f"C{sel_row}")


def _scenario_cells(ws):
    ws.cell(row=sel_row, column=1, value="Active scenario (1 = Base, 2 = Bull, 3 = Bear)").font = BOLD
    c = ws.cell(row=sel_row, column=3, value=A.get("scenario", 1))
    c.font, c.fill = BLUE, KEY_FILL
    d = ws.cell(row=sel_row, column=4, value=f'=CHOOSE(C{sel_row},"Base","Bull","Bear")')
    d.font = BOLD


asb.items[-1] = ("custom", sel_row, _scenario_cells)
basis_row = sel_row + 1
SC["ret_basis"] = ("Assumptions", f"C{basis_row}")


def _basis_cells(ws):
    ws.cell(row=basis_row, column=1, value="Shareholder returns based on (1 = % of net income, 2 = % of FCF)").font = BOLD
    c = ws.cell(row=basis_row, column=3, value=A.get("return_basis", 1))
    c.font, c.fill = BLUE, KEY_FILL
    ws.cell(row=basis_row, column=4, value=f'=CHOOSE(C{basis_row},"net income","free cash flow")').font = BLACK


asb.items.append(("custom", basis_row, _basis_cells))
asb.r += 3
fc_idx = lambda j: j - NH  # noqa: E731


def inp(series):
    return lambda j: series[fc_idx(j)]


def live(base_key, bull_key, bear_key):
    return lambda j: (f"=CHOOSE({sref('scenario')},{ref('Assumptions', base_key, j)},"
                      f"{ref('Assumptions', bull_key, j)},{ref('Assumptions', bear_key, j)})")


DRV = M.get("drivers") or {}
DRIVERS = "Drivers"
if DRV:
    # Revenue comes from the Drivers tab; the scenario totals below are the draft, shown for reference.
    asb.hdr("Revenue growth - set on the Drivers tab (segment volume and price); draft totals shown as memo")
    asb.line("g_base", "  Base (memo: draft total)", PCT, fc=lambda j: f"={A['base']['revenue_growth'][fc_idx(j)]}",
             note="Memo only. Change revenue on the Drivers tab: segment volume and price growth by scenario.")
    asb.line("g_bull", "  Bull (memo: draft total)", PCT, fc=lambda j: f"={A['bull']['revenue_growth'][fc_idx(j)]}")
    asb.line("g_bear", "  Bear (memo: draft total)", PCT, fc=lambda j: f"={A['bear']['revenue_growth'][fc_idx(j)]}")
    asb.line("g_live", "Revenue growth - active scenario (from Drivers)", PCT, bold=True,
             hist=lambda j: None if j == 0 else f"={ref('Income Statement', 'growth', j)}",
             fc=lambda j: f"={ref('Income Statement', 'growth', j)}")
else:
    asb.hdr("Revenue growth")
    asb.line("g_base", "  Base", PCT, fc=inp(A["base"]["revenue_growth"]), key_fill=True,
             note="Consensus-anchored years use Distilla consensus mean. " + " ".join(M.get("notes", [])))
    asb.line("g_bull", "  Bull", PCT, fc=inp(A["bull"]["revenue_growth"]), note="Consensus high estimates")
    asb.line("g_bear", "  Bear", PCT, fc=inp(A["bear"]["revenue_growth"]), note="Consensus low estimates")
    asb.line("g_live", "Revenue growth - active scenario", PCT, bold=True,
             hist=lambda j: None if j == 0 else f"={ref('Income Statement', 'growth', j)}",
             fc=live("g_base", "g_bull", "g_bear"))
DMG = bool(DRV.get("margins_on"))
if DMG:
    # EBIT comes from segment margins on the Drivers tab; the company margins below are the draft result (memo).
    asb.hdr("EBIT margin - set on the Drivers tab (segment margins); draft company result shown as memo")
    asb.line("m_base", "  Base (memo: draft mix)", PCT, fc=lambda j: f"={A['base']['ebit_margin'][fc_idx(j)]}",
             note="Memo only. Change margins on the Drivers tab: segment margins by scenario and the corporate line.")
    asb.line("m_bull", "  Bull (memo: draft mix)", PCT, fc=lambda j: f"={A['bull']['ebit_margin'][fc_idx(j)]}")
    asb.line("m_bear", "  Bear (memo: draft mix)", PCT, fc=lambda j: f"={A['bear']['ebit_margin'][fc_idx(j)]}")
    asb.line("m_live", "EBIT margin - active scenario (from Drivers)", PCT, bold=True,
             hist=lambda j: f"={ref('Income Statement', 'ebit_m', j)}", fc=lambda j: f"={ref('Income Statement', 'ebit_m', j)}")
else:
    asb.hdr("EBIT margin")
    asb.line("m_base", "  Base", PCT, fc=inp(A["base"]["ebit_margin"]), key_fill=True)
    asb.line("m_bull", "  Bull", PCT, fc=inp(A["bull"]["ebit_margin"]))
    asb.line("m_bear", "  Bear", PCT, fc=inp(A["bear"]["ebit_margin"]))
    asb.line("m_live", "EBIT margin - active scenario", PCT, bold=True,
             hist=lambda j: f"={ref('Income Statement', 'ebit_m', j)}", fc=live("m_base", "m_bull", "m_bear"))
RC = M.get("recurring_charges") or {}
if RC.get("pct_rev"):
    asb.line("rc_memo", "  memo: basis gap already taken off (company basis -> Distilla EBIT basis)", PCT,
             fc=lambda j: RC["pct_rev"],
             note="Consensus follows the company's own operating profit; the median gap to Distilla's EBIT is taken off "
                  "every consensus-derived margin above. Source: " + (RC.get("source") or "not recorded"))

asb.hdr("Operating drivers (history shown for reference)")


def hist_ratio(expr):
    return lambda j: f"=IFERROR({expr(j)},0)"


asb.line("opex_pct", "Operating expenses % revenue (SG&A, R&D, other)", PCT,
         hist=hist_ratio(lambda j: f"{ref('Income Statement', 'opex', j)}/{ref('Income Statement', 'rev', j)}"),
         fc=inp(A["opex_pct_rev"]),
         note="Consensus (gross profit - EBIT) / sales where available, then held. Gross margin = EBIT margin + this, so COGS follows the scenario.")
asb.line("da_ovr", "D&A % revenue - consensus years (blank = use depreciation rate)", PCT,
         hist=hist_ratio(lambda j: f"{ref('Income Statement', 'da', j)}/{ref('Income Statement', 'rev', j)}"),
         fc=lambda j: A["da_override_pct_rev"][fc_idx(j)],
         note="Consensus EBITDA - EBIT, divided by consensus sales. Leave blank to depreciate off the asset base.")
asb.line("dep_rate", "Depreciation rate on beginning net PP&E", PCT,
         hist=lambda j: None if j == 0 else f"=IFERROR({ref('Income Statement', 'da', j)}/{ref('Balance Sheet', 'ppe', j - 1)},0)",
         fc=inp(A["dep_rate"]), note="Median of recent years. Used whenever the consensus D&A cell is blank.")
asb.line("capex_pct", "Capex % revenue", PCT,
         hist=hist_ratio(lambda j: f"-{raw('cash_flow_capital_expenditures', j)}/{ref('Income Statement', 'rev', j)}"),
         fc=inp(A["capex_pct_rev"]), key_fill=True,
         note="Consensus capex / sales where available. Afterwards the level that keeps PP&E / revenue stable: ratio x (dep. rate + growth) / (1 + growth).")
asb.line("dso", "Receivable days (on revenue)", DAYS,
         hist=hist_ratio(lambda j: f"{ref('Balance Sheet', 'rec', j)}/{ref('Income Statement', 'rev', j)}*365"),
         fc=inp(A["dso"]))
asb.line("dio", "Inventory days (on COGS)", DAYS,
         hist=hist_ratio(lambda j: f"{ref('Balance Sheet', 'inv', j)}/{ref('Income Statement', 'cogs', j)}*365"),
         fc=inp(A["dio"]))
asb.line("dpo", "Payable days (on COGS)", DAYS,
         hist=hist_ratio(lambda j: f"{ref('Balance Sheet', 'ap', j)}/{ref('Income Statement', 'cogs', j)}*365"),
         fc=inp(A["dpo"]))
asb.line("oca_pct", "Other current assets % revenue", PCT,
         hist=hist_ratio(lambda j: f"{ref('Balance Sheet', 'oca', j)}/{ref('Income Statement', 'rev', j)}"),
         fc=inp(A["oca_pct_rev"]))
asb.line("ocl_pct", "Other current liabilities % revenue", PCT,
         hist=hist_ratio(lambda j: f"{ref('Balance Sheet', 'ocl', j)}/{ref('Income Statement', 'rev', j)}"),
         fc=inp(A["ocl_pct_rev"]))
asb.line("lease_add", "New leases (right-of-use additions) % revenue", PCT,
         hist=hist_ratio(lambda j: f"-{raw('cash_flow_repayments_of_operating_lease_liabilities', j)}/{ref('Income Statement', 'rev', j)}"),
         fc=inp(A.get("lease_add_pct_rev", [0] * NF)),
         note="IFRS 16 lessees: new leases treated as debt-financed asset additions (history shows principal repaid / revenue). 0 under US GAAP.")
asb.line("lease_rate", "Lease principal repaid % of opening lease liability", PCT,
         hist=lambda j: None if j == 0 else f"=IFERROR(-{raw('cash_flow_repayments_of_operating_lease_liabilities', j)}/{ref('Balance Sheet', 'lease', j - 1)},0)",
         fc=inp(A.get("lease_repay_rate", [0] * NF)))
asb.line("tax", "Effective tax rate", PCT, hist=lambda j: f"={ref('Income Statement', 'etr', j)}",
         fc=inp(A["tax_rate"]))
asb.line("cash_yield", "Interest yield on cash", PCT, fc=inp(A["cash_yield"]),
         note="Applied to beginning-of-year cash (avoids circularity).")
asb.line("kd", "Pre-tax cost of debt", PCT, fc=inp(A["cost_of_debt"]),
         note="Applied to beginning-of-year debt (avoids circularity).")
asb.line("payout", "Dividends (% of net income or FCF - see basis above)", PCT,
         hist=hist_ratio(lambda j: f"-{raw('cash_flow_cash_dividends_paid', j)}/{ref('Income Statement', 'ni', j)}"),
         fc=inp(A["payout"]))
asb.line("bb_pct", "Buybacks (% of net income or FCF - see basis above)", PCT,
         hist=hist_ratio(lambda j: f"-{raw('cash_flow_repurchase_of_common_and_preferred_stock', j)}/{ref('Income Statement', 'ni', j)}"),
         fc=inp(A["buyback_pct_ni"]), note="Buybacks are capped at available cash (Schedules); Checks counts the capped years.")
asb.line("debt_iss", f"Net debt issuance / (repayment)", NUM,
         hist=lambda j: f"={raw('cash_flow_issuance_or_reduction_of_debt_net', j)}", fc=inp(A["net_debt_issuance"]),
         unit=f"{CUR} {UNITS}")
asb.line("mi_pct", "Minority share of consolidated NI", PCT,
         hist=hist_ratio(lambda j: f"-{ref('Income Statement', 'mi', j)}/{ref('Income Statement', 'cons_ni', j)}"),
         fc=inp(A["mi_pct"]))

# consensus reference block
CP = M["consensus"]["periods"]
CS = M["consensus"]["series"]
asb.hdr("Consensus reference - Distilla latest snapshot (memo, not linked)")
for key, lab in (("sales_mean", "Sales - mean"), ("sales_high", "Sales - high"), ("sales_low", "Sales - low"),
                 ("gross_inc_mean", "Gross profit - mean"),
                 ("ebit_mean", "EBIT - mean"), ("ebit_high", "EBIT - high"), ("ebit_low", "EBIT - low"),
                 ("ebitda_mean", "EBITDA - mean"), ("capex_mean", "Capex - mean")):
    if key not in CS:
        continue
    vals = {p: v for p, v in zip(CP, CS[key])}
    asb.line("c_" + key, lab, NUM, fc=(lambda vals: lambda j: vals.get(FP[fc_idx(j)]))(vals))
if "sales_mean" in CS:
    asb.line("c_var", "Model revenue vs consensus mean", PCT,
             fc=lambda j: (f"=IFERROR({ref('Income Statement', 'rev', j)}/{ref('Assumptions', 'c_sales_mean', j)}-1,0)"
                           if FP[fc_idx(j)] in CP else None))

# ---------------------------------------------------------------- Drivers tab
# Segment revenue = last year x (1 + volume) x (1 + price), or units x revenue per unit where a unit series
# exists. Each driver has Base / Bull / Bear inputs and an active-scenario line; "Other / eliminations"
# (intersegment sales, unallocated revenue) grows with the segment total. The total feeds the Model tab.
if DRV:
    drv = SB(DRIVERS, "Revenue drivers - segment volume x price (blue = input)")
    SEGS, DH, DU = DRV["segments"], DRV["hist"], DRV.get("units") or {}
    drv.text(f"Revenue lines: {DRV.get('basis_type', 'business segments')}. Source: " + DRV.get("source", "")
             + (f" - field '{DRV['field']}'" if DRV.get("field") else ""))
    drv.text("Draft: years with consensus (and revenue anchors) are calibrated so the segments sum to the scenario "
             "total; later years run on the drivers. Change any blue cell - revenue follows.")

    def d_live(stem):
        return lambda j: (f"=CHOOSE({sref('scenario')},{ref(DRIVERS, stem + '_base', j)},"
                          f"{ref(DRIVERS, stem + '_bull', j)},{ref(DRIVERS, stem + '_bear', j)})")

    for k, s in enumerate(SEGS):
        key = f"s{k}"
        drv.hdr(s + (f"  - {DRV['basis'][s]}" if DRV.get("basis", {}).get(s) else "")
                + (f"  [history: {DRV['line_sources'][s]}]" if DRV.get("line_sources", {}).get(s) else ""))
        for kind, lab in (("market", "memo: market growth (volume = (1 + market) x (1 + share) - 1)"),
                          ("share", "memo: share change")):
            for sc in ("base", "bull", "bear"):
                vals = ((DRV.get(kind) or {}).get(sc) or {}).get(s) or []
                if any(v is not None for v in vals):
                    drv.line(f"{key}_{kind}_{sc}", f"  {lab} - {sc.title()}", PCT,
                             fc=(lambda vals: lambda j: vals[fc_idx(j)])(vals))
        for kind, lab in (("v", "Volume growth"), ("p", "Price growth")):
            for sc in ("base", "bull", "bear"):
                vals = DRV["scen"][sc][s]["volume" if kind == "v" else "price"]
                drv.line(f"{key}_{kind}_{sc}", f"  {lab} - {sc.title()}", PCT,
                         fc=(lambda vals: lambda j: vals[fc_idx(j)])(vals), key_fill=(sc == "base"))
            drv.line(f"{key}_{kind}", f"{lab} - active scenario", PCT, bold=True, fc=d_live(f"{key}_{kind}"))
        hv = DH[s]
        if s in DU:
            uv = DU[s]
            drv.line(f"{key}_u", f"Units ({DRV.get('units_unit') or 'as reported'})", NUM1,
                     hist=(lambda uv: lambda j: uv[j])(uv),
                     fc=(lambda key: lambda j: f"={ref(DRIVERS, key + '_u', prev(j))}*(1+{ref(DRIVERS, key + '_v', j)})")(key))
            drv.line(f"{key}_asp", "Revenue per unit (revenue / units, ‡)", PS,
                     hist=(lambda key, uv: lambda j: (f"=IFERROR({ref(DRIVERS, key + '_r', j)}/{ref(DRIVERS, key + '_u', j)},0)"
                                                      if uv[j] else None))(key, uv),
                     fc=(lambda key: lambda j: f"={ref(DRIVERS, key + '_asp', prev(j))}*(1+{ref(DRIVERS, key + '_p', j)})")(key))
            rfc = (lambda key: lambda j: f"={ref(DRIVERS, key + '_u', j)}*{ref(DRIVERS, key + '_asp', j)}")(key)
        else:
            rfc = (lambda key: lambda j: (f"={ref(DRIVERS, key + '_r', prev(j))}*(1+{ref(DRIVERS, key + '_v', j)})"
                                          f"*(1+{ref(DRIVERS, key + '_p', j)})"))(key)
        drv.line(f"{key}_r", f"Revenue - {s}", NUM, bold=True, hist=(lambda hv: lambda j: hv[j])(hv), fc=rfc)
        drv.line(f"{key}_g", "  growth %", PCT,
                 hist=(lambda key, hv: lambda j: (f"=IFERROR({ref(DRIVERS, key + '_r', j)}/{ref(DRIVERS, key + '_r', prev(j))}-1,0)"
                                                  if j > 0 and hv[j] is not None and hv[j - 1] is not None else None))(key, hv),
                 fc=(lambda key: lambda j: f"=IFERROR({ref(DRIVERS, key + '_r', j)}/{ref(DRIVERS, key + '_r', prev(j))}-1,0)")(key))
        if DMG:
            for sc in ("base", "bull", "bear"):
                vals = DRV["margin"][sc][s]
                drv.line(f"{key}_m_{sc}", f"  Profit margin - {sc.title()}", PCT,
                         fc=(lambda vals: lambda j: vals[fc_idx(j)])(vals), key_fill=(sc == "base"))
            ph = DRV["profit_hist"][s]
            drv.line(f"{key}_m", "Profit margin - active scenario", PCT, bold=True,
                     hist=(lambda key, ph: lambda j: (f"=IFERROR({ref(DRIVERS, key + '_op', j)}/{ref(DRIVERS, key + '_r', j)},0)"
                                                      if ph[j] is not None else None))(key, ph),
                     fc=d_live(f"{key}_m"))
            drv.line(f"{key}_op", f"Segment profit - {s}", NUM, bold=True, hist=(lambda ph: lambda j: ph[j])(ph),
                     fc=(lambda key: lambda j: f"={ref(DRIVERS, key + '_r', j)}*{ref(DRIVERS, key + '_m', j)}")(key))
    _complete = [all(DH[s][j] is not None for s in SEGS) for j in range(NH)]
    seg_sum_f = lambda j: "=" + "+".join(ref(DRIVERS, f"s{k}_r", j) for k in range(len(SEGS)))  # noqa: E731
    drv.hdr("Total revenue")
    drv.line("seg_sum", "Segment total", NUM, bold=True, hist=lambda j: seg_sum_f(j) if _complete[j] else None, fc=seg_sum_f)
    drv.line("other", "Other / eliminations (Distilla revenue - segment total)", NUM,
             hist=lambda j: f"={raw('income_statement_sales', j)}-{ref(DRIVERS, 'seg_sum', j)}" if _complete[j] else None,
             fc=lambda j: (f"={ref(DRIVERS, 'other', prev(j))}*IFERROR({ref(DRIVERS, 'seg_sum', j)}"
                           f"/{ref(DRIVERS, 'seg_sum', prev(j))},1)"),
             note="Intersegment sales and unallocated revenue. Forecast: grows with the segment total.")
    drv.line("total", "Revenue (feeds the Model tab)", NUM, bold=True, top=True, key_fill=True,
             hist=lambda j: f"={raw('income_statement_sales', j)}",
             fc=lambda j: f"={ref(DRIVERS, 'seg_sum', j)}+{ref(DRIVERS, 'other', j)}")
    drv.line("total_g", "  growth %", PCT, hist=lambda j: None if j == 0 else f"=IFERROR({ref(DRIVERS, 'total', j)}/{ref(DRIVERS, 'total', prev(j))}-1,0)",
             fc=lambda j: f"=IFERROR({ref(DRIVERS, 'total', j)}/{ref(DRIVERS, 'total', prev(j))}-1,0)")
    if "sales_mean" in M["consensus"]["series"]:
        _cp = M["consensus"]["periods"]
        drv.line("vs_cons", "Revenue vs consensus mean", PCT,
                 fc=lambda j: (f"=IFERROR({ref(DRIVERS, 'total', j)}/{ref('Assumptions', 'c_sales_mean', j)}-1,0)"
                               if FP[fc_idx(j)] in _cp else None))
    if DMG:
        # EBIT = segment profit + corporate / unallocated (Distilla EBIT - segment profit: corporate costs,
        # restructuring, the basis gap between segment and consolidated profit), held as a % of revenue.
        _pc = [all(DRV["profit_hist"][s][j] is not None for s in SEGS) for j in range(NH)]
        op_sum_f = lambda j: "=" + "+".join(ref(DRIVERS, f"s{k}_op", j) for k in range(len(SEGS)))  # noqa: E731
        drv.hdr("EBIT - segment margins (feeds the Model tab)")
        drv.text("Draft: in consensus years the segment margins are calibrated so EBIT matches consensus EBIT (on "
                 "Distilla's basis); later years follow each line's anchor, else the company-level path. Change any "
                 "blue margin - EBIT follows.")
        drv.line("op_sum", "Segment profit total", NUM, bold=True, hist=lambda j: op_sum_f(j) if _pc[j] else None, fc=op_sum_f)
        drv.line("corp_pct", "Corporate / unallocated % of revenue", PCT,
                 hist=lambda j: f"=IFERROR({ref(DRIVERS, 'corp', j)}/{ref(DRIVERS, 'total', j)},0)" if _pc[j] else None,
                 fc=(lambda vals: lambda j: vals[fc_idx(j)])(DRV["corp_pct"]), key_fill=True,
                 note="Held at the last actual year's share (all scenarios).")
        drv.line("corp", "Corporate / unallocated (Distilla EBIT - segment profit)", NUM,
                 hist=lambda j: (f"={raw('income_statement_ebit_operating_income', j)}-{ref(DRIVERS, 'op_sum', j)}"
                                 if _pc[j] else None),
                 fc=lambda j: f"={ref(DRIVERS, 'total', j)}*{ref(DRIVERS, 'corp_pct', j)}",
                 note="Corporate costs, restructuring and other items not in segment profit.")
        drv.line("ebit", "EBIT (feeds the Model tab)", NUM, bold=True, top=True, key_fill=True,
                 hist=lambda j: f"={raw('income_statement_ebit_operating_income', j)}",
                 fc=lambda j: f"={ref(DRIVERS, 'op_sum', j)}+{ref(DRIVERS, 'corp', j)}")
        drv.line("ebit_m", "  EBIT margin %", PCT, hist=lambda j: f"=IFERROR({ref(DRIVERS, 'ebit', j)}/{ref(DRIVERS, 'total', j)},0)",
                 fc=lambda j: f"=IFERROR({ref(DRIVERS, 'ebit', j)}/{ref(DRIVERS, 'total', j)},0)")
        for sc in ("base", "bull", "bear"):
            drv.line(f"cm_{sc}", f"  memo: company-level margin path - {sc.title()} (cross-check)", PCT,
                     fc=(lambda vals: lambda j: vals[fc_idx(j)])(DRV["company_path"][sc]))

# ---------------------------------------------------------------- Model tab
# Four logical statements share one physical sheet. History links to Raw Data line by line;
# every subtotal is a real sum, and any gap to Distilla's reported subtotal is shown on an
# explicit "Unreconciled vs Distilla" line (expected to be 0) instead of being hidden in a plug.
model = SB(MODEL, "3-Statement Model - Income Statement, Balance Sheet, Cash Flow, Schedules")
IS, BS, CF, SCH = "Income Statement", "Balance Sheet", "Cash Flow", "Schedules"
prev = lambda j: j - 1  # noqa: E731
flat = lambda sheet, key: (lambda j: f"={ref(sheet, key, prev(j))}")  # noqa: E731
rsum = lambda sheet, k0, k1: (lambda j: f"=SUM({col(j)}{ROWS[(sheet, k0)]}:{col(j)}{ROWS[(sheet, k1)]})")  # noqa: E731
link = lambda m: (lambda j: f"={raw(m, j)}")  # noqa: E731

# ---------- Income statement
isb = Section(model, IS, "INCOME STATEMENT")
isb.line("rev", "Revenue", NUM, bold=True, hist=link('income_statement_sales'),
         fc=(lambda j: f"={ref(DRIVERS, 'total', j)}") if DRV else
         (lambda j: f"={ref(IS, 'rev', prev(j))}*(1+{ref('Assumptions', 'g_live', j)})"))
isb.line("growth", "  growth %", PCT, hist=lambda j: None if j == 0 else f"=IFERROR({ref(IS, 'rev', j)}/{ref(IS, 'rev', prev(j))}-1,0)",
         fc=lambda j: f"=IFERROR({ref(IS, 'rev', j)}/{ref(IS, 'rev', prev(j))}-1,0)")
isb.line("cogs", "Cost of goods sold (incl. D&A)", NUM, hist=link('income_statement_cost_of_goods_sold_cogs_incl_d_and_a'),
         fc=lambda j: f"={ref(IS, 'rev', j)}-{ref(IS, 'gp', j)}")
isb.line("gp", "Gross profit", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(IS, 'rev', j)}-{ref(IS, 'cogs', j)}", fc=lambda j: f"={ref(IS, 'ebit', j)}+{ref(IS, 'opex', j)}",
         note="Forecast: EBIT + operating expenses, so gross margin moves with the scenario and opex cannot go negative.")
isb.line("gm", "  gross margin %", PCT, hist=lambda j: f"=IFERROR({ref(IS, 'gp', j)}/{ref(IS, 'rev', j)},0)",
         fc=lambda j: f"=IFERROR({ref(IS, 'gp', j)}/{ref(IS, 'rev', j)},0)")
isb.line("opex", "Operating expenses (SG&A, R&D, other)", NUM,
         hist=lambda j: f"={ref(IS, 'gp', j)}-{ref(IS, 'ebit', j)}", fc=lambda j: f"={ref(IS, 'rev', j)}*{ref('Assumptions', 'opex_pct', j)}",
         note="History: gross profit - EBIT (Distilla SG&A already includes R&D). Forecast: revenue x opex %.")
isb.line("ebit", "EBIT (operating income)", NUM, bold=True, top=True, hist=link('income_statement_ebit_operating_income'),
         fc=(lambda j: f"={ref(DRIVERS, 'ebit', j)}") if DMG else
         (lambda j: f"={ref(IS, 'rev', j)}*{ref('Assumptions', 'm_live', j)}"))
isb.line("ebit_m", "  EBIT margin %", PCT, hist=lambda j: f"=IFERROR({ref(IS, 'ebit', j)}/{ref(IS, 'rev', j)},0)",
         fc=lambda j: f"=IFERROR({ref(IS, 'ebit', j)}/{ref(IS, 'rev', j)},0)")
isb.line("da", "D&A (memo)", NUM, hist=link('cash_flow_depreciation_depletion_and_amortization'),
         fc=lambda j: f"={ref(SCH, 'da', j)}")
isb.line("ebitda", "EBITDA", NUM, bold=True, hist=lambda j: f"={ref(IS, 'ebit', j)}+{ref(IS, 'da', j)}",
         fc=lambda j: f"={ref(IS, 'ebit', j)}+{ref(IS, 'da', j)}")
isb.blank()
isb.line("int_inc", "Interest income", NUM, fc=lambda j: f"={ref(SCH, 'int_inc', j)}")
isb.line("int_exp", "Interest expense", NUM, fc=lambda j: f"=-{ref(SCH, 'int_exp', j)}")
isb.line("nonop", "Net non-operating income / (expense)", NUM,
         hist=lambda j: f"={raw('income_statement_pretax_income', j)}-{ref(IS, 'ebit', j)}",
         fc=lambda j: f"={ref(IS, 'int_inc', j)}+{ref(IS, 'int_exp', j)}",
         note="History: pretax income minus EBIT (interest, FX, one-offs; Distilla interest lines are often blank). Forecast: interest only.")
isb.line("pretax", "Pretax income", NUM, bold=True, top=True, hist=link('income_statement_pretax_income'),
         fc=lambda j: f"={ref(IS, 'ebit', j)}+{ref(IS, 'nonop', j)}")
isb.line("taxes", "Income taxes", NUM, hist=lambda j: f"=-{raw('income_statement_income_taxes', j)}",
         fc=lambda j: f"=-{ref(IS, 'pretax', j)}*{ref('Assumptions', 'tax', j)}")
isb.line("etr", "  effective tax rate %", PCT, hist=lambda j: f"=IFERROR(-{ref(IS, 'taxes', j)}/{ref(IS, 'pretax', j)},0)",
         fc=lambda j: f"=IFERROR(-{ref(IS, 'taxes', j)}/{ref(IS, 'pretax', j)},0)")
isb.line("other_at", "Other after-tax items (affiliates, disc. ops)", NUM,
         hist=lambda j: (f"={raw('income_statement_net_income', j)}+{raw('income_statement_minority_interest', j)}"
                         f"-({ref(IS, 'pretax', j)}+{ref(IS, 'taxes', j)})"),
         note="History: equity-method income and discontinued items, derived so net income ties to Distilla. Not forecast.")
isb.line("cons_ni", "Consolidated net income", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(IS, 'pretax', j)}+{ref(IS, 'taxes', j)}+{ref(IS, 'other_at', j)}",
         fc=lambda j: f"={ref(IS, 'pretax', j)}+{ref(IS, 'taxes', j)}+{ref(IS, 'other_at', j)}")
isb.line("mi", "Minority interest", NUM, hist=lambda j: f"=-{raw('income_statement_minority_interest', j)}",
         fc=lambda j: f"=-{ref(IS, 'cons_ni', j)}*{ref('Assumptions', 'mi_pct', j)}")
isb.line("ni", "Net income to common", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(IS, 'cons_ni', j)}+{ref(IS, 'mi', j)}", fc=lambda j: f"={ref(IS, 'cons_ni', j)}+{ref(IS, 'mi', j)}")
isb.blank()
isb.line("shares", "Diluted shares", NUM1, unit="m", hist=link('income_statement_diluted_shares_outstanding'),
         fc=lambda j: f"=MAX(0,{ref(IS, 'shares', prev(j))}-{ref(SCH, 'bb', j)}/({sref('price')}*{sref('fx')}))",
         note="Forecast: prior shares less buybacks at today's price (EPS only; valuation uses current diluted shares).")
isb.line("eps", "Diluted EPS", PS, unit=CUR, hist=lambda j: f"=IFERROR({ref(IS, 'ni', j)}/{ref(IS, 'shares', j)},0)",
         fc=lambda j: f"=IFERROR({ref(IS, 'ni', j)}/{ref(IS, 'shares', j)},0)")
isb.line("dps", "Dividends per share", PS, unit=CUR, hist=link('income_statement_dividends_per_share'),
         fc=lambda j: f"=IFERROR({ref(SCH, 'div', j)}/{ref(IS, 'shares', j)},0)")
isb.close()

# ---------- Balance sheet
bsb = Section(model, BS, "BALANCE SHEET")
bsb.hdr("Assets")
bsb.line("cash", "Cash & ST investments", NUM, hist=link('balance_sheet_cash_and_short_term_investments'),
         fc=lambda j: f"={ref(CF, 'cash_end', j)}", note="Forecast cash comes from the cash flow statement (balancing item).")
bsb.line("rec", "Receivables", NUM, hist=link('balance_sheet_short_term_receivables'), fc=lambda j: f"={ref(SCH, 'rec', j)}")
bsb.line("inv", "Inventories", NUM, hist=link('balance_sheet_inventories'), fc=lambda j: f"={ref(SCH, 'inv', j)}")
bsb.line("oca", "Other current assets", NUM, hist=link('balance_sheet_other_current_assets'), fc=lambda j: f"={ref(SCH, 'oca', j)}")
bsb.line("u_ca", "  Unreconciled vs Distilla (current assets)", NUM,
         hist=lambda j: f"={raw('balance_sheet_total_current_assets', j)}-SUM({col(j)}{ROWS[(BS, 'cash')]}:{col(j)}{ROWS[(BS, 'oca')]})",
         fc=flat(BS, "u_ca"), note="Distilla total minus the itemised lines. Expected 0; any amount is shown, not hidden.")
bsb.line("tca", "Total current assets", NUM, bold=True, top=True, hist=rsum(BS, "cash", "u_ca"), fc=rsum(BS, "cash", "u_ca"))
bsb.line("ppe", "Net PP&E", NUM, hist=link('balance_sheet_net_property_plant_and_equipment'), fc=lambda j: f"={ref(SCH, 'ppe_end', j)}")
bsb.line("lti", "Long-term investments", NUM, hist=link('balance_sheet_total_long_term_investments'), fc=flat(BS, "lti"))
bsb.line("intang", "Intangibles incl. goodwill", NUM, hist=link('balance_sheet_intangible_assets'), fc=flat(BS, "intang"),
         note="Held flat; all forecast D&A runs through PP&E (simplification).")
bsb.line("dta", "Deferred tax assets", NUM, hist=link('balance_sheet_deferred_tax_assets'), fc=flat(BS, "dta"))
bsb.line("olta", "Other long-term assets", NUM, hist=link('balance_sheet_other_assets'), fc=flat(BS, "olta"))
bsb.line("u_lta", "  Unreconciled vs Distilla (long-term assets)", NUM,
         hist=lambda j: (f"={raw('balance_sheet_total_assets', j)}-{ref(BS, 'tca', j)}"
                         f"-SUM({col(j)}{ROWS[(BS, 'ppe')]}:{col(j)}{ROWS[(BS, 'olta')]})"),
         fc=flat(BS, "u_lta"), note="Expected 0.")
bsb.line("ta", "Total assets", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(BS, 'tca', j)}+SUM({col(j)}{ROWS[(BS, 'ppe')]}:{col(j)}{ROWS[(BS, 'u_lta')]})",
         fc=lambda j: f"={ref(BS, 'tca', j)}+SUM({col(j)}{ROWS[(BS, 'ppe')]}:{col(j)}{ROWS[(BS, 'u_lta')]})")
bsb.hdr("Liabilities & equity")
bsb.line("std", "Short-term debt", NUM, hist=link('balance_sheet_short_term_debt_and_curr_portion_long_term_debt'), fc=flat(BS, "std"))
bsb.line("ap", "Accounts payable", NUM, hist=link('balance_sheet_accounts_payable'), fc=lambda j: f"={ref(SCH, 'ap', j)}")
bsb.line("taxp", "Income tax payable", NUM, hist=link('balance_sheet_income_tax_payable'), fc=flat(BS, "taxp"))
bsb.line("ocl", "Other current liabilities", NUM, hist=link('balance_sheet_other_current_liabilities'), fc=lambda j: f"={ref(SCH, 'ocl', j)}")
bsb.line("u_cl", "  Unreconciled vs Distilla (current liabilities)", NUM,
         hist=lambda j: f"={raw('balance_sheet_total_current_liabilities', j)}-SUM({col(j)}{ROWS[(BS, 'std')]}:{col(j)}{ROWS[(BS, 'ocl')]})",
         fc=flat(BS, "u_cl"), note="Expected 0.")
bsb.line("tcl", "Total current liabilities", NUM, bold=True, top=True, hist=rsum(BS, "std", "u_cl"), fc=rsum(BS, "std", "u_cl"))
bsb.line("ltd", "Long-term debt (excl. leases)", NUM, hist=link('balance_sheet_long_term_debt_excl_lease_obligations'),
         fc=lambda j: f"={ref(SCH, 'debt_end', j)}-{ref(BS, 'std', j)}")
bsb.line("lease", "Lease liabilities", NUM, hist=link('balance_sheet_capital_and_operating_lease_obligations'), fc=lambda j: f"={ref(SCH, 'lease_end', j)}")
bsb.line("prov", "Provisions", NUM, hist=link('balance_sheet_provision_for_risks_and_charges'), fc=flat(BS, "prov"))
bsb.line("dtl", "Deferred tax liabilities", NUM, hist=link('balance_sheet_deferred_tax_liabilities'), fc=flat(BS, "dtl"))
bsb.line("oltl", "Other long-term liabilities", NUM, hist=link('balance_sheet_other_liabilities'), fc=flat(BS, "oltl"))
bsb.line("u_ltl", "  Unreconciled vs Distilla (long-term liabilities)", NUM,
         hist=lambda j: (f"={raw('balance_sheet_total_liabilities', j)}-{ref(BS, 'tcl', j)}"
                         f"-SUM({col(j)}{ROWS[(BS, 'ltd')]}:{col(j)}{ROWS[(BS, 'oltl')]})"),
         fc=flat(BS, "u_ltl"), note="Expected 0.")
bsb.line("tl", "Total liabilities", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(BS, 'tcl', j)}+SUM({col(j)}{ROWS[(BS, 'ltd')]}:{col(j)}{ROWS[(BS, 'u_ltl')]})",
         fc=lambda j: f"={ref(BS, 'tcl', j)}+SUM({col(j)}{ROWS[(BS, 'ltd')]}:{col(j)}{ROWS[(BS, 'u_ltl')]})")
bsb.line("se", "Shareholders' equity", NUM, hist=link('balance_sheet_total_shareholders_equity'), fc=lambda j: f"={ref(SCH, 'eq_end', j)}")
bsb.line("mi_bs", "Minority interest", NUM, hist=link('balance_sheet_accumulated_minority_interest'), fc=lambda j: f"={ref(SCH, 'mi_end', j)}")
bsb.line("u_eq", "  Unreconciled vs Distilla (equity / other instruments)", NUM,
         hist=lambda j: f"={raw('balance_sheet_total_assets', j)}-{ref(BS, 'tl', j)}-{ref(BS, 'se', j)}-{ref(BS, 'mi_bs', j)}",
         fc=flat(BS, "u_eq"),
         note="Distilla total assets less liabilities, shareholders' equity and minority interest. Expected 0; a non-zero amount "
              "means Distilla's own balance sheet does not add up (e.g. perpetual bonds or other equity instruments not itemised).")
bsb.line("te", "Total equity", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(BS, 'se', j)}+{ref(BS, 'mi_bs', j)}+{ref(BS, 'u_eq', j)}",
         fc=lambda j: f"={ref(BS, 'se', j)}+{ref(BS, 'mi_bs', j)}+{ref(BS, 'u_eq', j)}")
bsb.line("tle", "Total liabilities & equity", NUM, bold=True,
         hist=lambda j: f"={ref(BS, 'tl', j)}+{ref(BS, 'te', j)}", fc=lambda j: f"={ref(BS, 'tl', j)}+{ref(BS, 'te', j)}")
bsb.blank()
bsb.line("check", "Balance check (assets - L&E)", NUM, bold=True,
         hist=lambda j: f"=ROUND({ref(BS, 'ta', j)}-{ref(BS, 'tle', j)},3)", fc=lambda j: f"=ROUND({ref(BS, 'ta', j)}-{ref(BS, 'tle', j)},3)")
bsb.close()

# ---------- Cash flow statement (history: every Distilla line; forecast: modelled lines only)
csb = Section(model, CF, "CASH FLOW STATEMENT")
csb.hdr("Operating activities")
csb.line("cons_ni", "Consolidated net income", NUM, hist=lambda j: f"={ref(IS, 'cons_ni', j)}", fc=lambda j: f"={ref(IS, 'cons_ni', j)}")
csb.line("da", "D&A", NUM, hist=lambda j: f"={ref(IS, 'da', j)}", fc=lambda j: f"={ref(IS, 'da', j)}")
csb.line("dtax", "Deferred taxes", NUM, hist=link('cash_flow_deferred_taxes'), note="History only; forecast treats book tax as cash tax.")
csb.line("noncash", "Other non-cash items (incl. stock-based comp.)", NUM,
         hist=lambda j: (f"={raw('cash_flow_funds_from_operations', j)}-{ref(CF, 'cons_ni', j)}-{ref(CF, 'da', j)}"
                         f"-{ref(CF, 'dtax', j)}"),
         note="History: Distilla funds from operations less net income, D&A and deferred tax. Mostly SBC. Not forecast: SBC treated as a real cost.")
csb.line("wc", "Change in working capital", NUM, hist=link('cash_flow_changes_in_working_capital'),
         fc=lambda j: f"=-{ref(SCH, 'dnwc', j)}", note="History: as reported. Forecast: from the NWC schedule.")
csb.line("cfo", "Cash from operations", NUM, bold=True, top=True, hist=rsum(CF, "cons_ni", "wc"), fc=rsum(CF, "cons_ni", "wc"))
csb.hdr("Investing activities")
csb.line("capex", "Capex (PP&E and intangibles)", NUM, hist=link('cash_flow_capital_expenditures'), fc=lambda j: f"=-{ref(SCH, 'capex', j)}")
csb.line("acq", "Acquisitions", NUM, hist=link('cash_flow_net_assets_from_acquisitions'))
csb.line("disp", "Sale of fixed assets & businesses", NUM, hist=link('cash_flow_sale_of_fixed_assets_and_businesses'))
csb.line("invs", "Net purchase / sale of investments", NUM, hist=link('cash_flow_purchase_or_sale_of_investments'))
csb.line("oinv", "Other investing", NUM, hist=link('cash_flow_other_investing_funds'))
csb.line("cfi", "Cash from investing", NUM, bold=True, top=True, hist=rsum(CF, "capex", "oinv"), fc=rsum(CF, "capex", "oinv"))
csb.hdr("Financing activities")
csb.line("div", "Dividends", NUM, hist=link('cash_flow_cash_dividends_paid'), fc=lambda j: f"=-{ref(SCH, 'div', j)}")
csb.line("bb", "Share buybacks", NUM, hist=link('cash_flow_repurchase_of_common_and_preferred_stock'), fc=lambda j: f"=-{ref(SCH, 'bb', j)}")
csb.line("iss", "Share issuance", NUM, hist=link('cash_flow_sale_of_common_and_preferred_stock'))
csb.line("debt", "Net debt issuance / (repayment)", NUM, hist=link('cash_flow_issuance_or_reduction_of_debt_net'),
         fc=lambda j: f"={ref(SCH, 'debt_iss', j)}")
csb.line("lease_rep", "Lease principal repayments", NUM, hist=link('cash_flow_repayments_of_operating_lease_liabilities'),
         fc=lambda j: f"=-{ref(SCH, 'lease_rep', j)}",
         note="From the lease schedule. New leases are non-cash additions to PP&E and the lease liability; repayments are the cash leg.")
csb.line("ofin", "Other financing", NUM, hist=link('cash_flow_other_financing_funds'))
csb.line("cff", "Cash from financing", NUM, bold=True, top=True, hist=rsum(CF, "div", "ofin"), fc=rsum(CF, "div", "ofin"))
csb.line("fx", "FX effect on cash", NUM, hist=link('cash_flow_exchange_rate_effect'))
csb.line("net", "Net change in cash", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(CF, 'cfo', j)}+{ref(CF, 'cfi', j)}+{ref(CF, 'cff', j)}+{ref(CF, 'fx', j)}",
         fc=lambda j: f"={ref(CF, 'cfo', j)}+{ref(CF, 'cfi', j)}+{ref(CF, 'cff', j)}+{ref(CF, 'fx', j)}")
csb.blank()
csb.line("cash_beg", "Beginning cash & ST investments", NUM,
         hist=lambda j: None if j == 0 else f"={ref(BS, 'cash', prev(j))}", fc=lambda j: f"={ref(BS, 'cash', prev(j))}")
csb.line("reclass", "Cash definition difference (ST investments)", NUM,
         hist=lambda j: None if j == 0 else f"={ref(BS, 'cash', j)}-{ref(CF, 'cash_beg', j)}-{ref(CF, 'net', j)}",
         note="Distilla's cash flow covers cash & equivalents; the balance sheet line also includes short-term investments. "
              "History: the difference, shown explicitly. Forecast: none (one definition).")
csb.line("cash_end", "Ending cash & ST investments", NUM, bold=True,
         hist=lambda j: None if j == 0 else f"={ref(CF, 'cash_beg', j)}+{ref(CF, 'net', j)}+{ref(CF, 'reclass', j)}",
         fc=lambda j: f"={ref(CF, 'cash_beg', j)}+{ref(CF, 'net', j)}+{ref(CF, 'reclass', j)}")
csb.blank()
csb.line("fcf", "Free cash flow (CFO - all capex, incl. intangibles)", NUM,
         hist=lambda j: f"={ref(CF, 'cfo', j)}+{ref(CF, 'capex', j)}", fc=lambda j: f"={ref(CF, 'cfo', j)}+{ref(CF, 'capex', j)}")
csb.line("fcf_d", "  memo: Distilla FCF definition (PP&E capex only)", NUM,
         hist=lambda j: f"={ref(CF, 'cfo', j)}+{raw('cash_flow_capital_expenditures_fixed_assets', j)}")
csb.close()

# ---------- Supporting schedules
ssb = Section(model, SCH, "SUPPORTING SCHEDULES")
ssb.hdr("Net working capital")
ssb.line("rec", "Receivables", NUM, hist=lambda j: f"={ref(BS, 'rec', j)}", fc=lambda j: f"={ref(IS, 'rev', j)}*{ref('Assumptions', 'dso', j)}/365")
ssb.line("inv", "Inventories", NUM, hist=lambda j: f"={ref(BS, 'inv', j)}", fc=lambda j: f"={ref(IS, 'cogs', j)}*{ref('Assumptions', 'dio', j)}/365")
ssb.line("oca", "Other current assets", NUM, hist=lambda j: f"={ref(BS, 'oca', j)}", fc=lambda j: f"={ref(IS, 'rev', j)}*{ref('Assumptions', 'oca_pct', j)}")
ssb.line("ap", "Accounts payable", NUM, hist=lambda j: f"={ref(BS, 'ap', j)}", fc=lambda j: f"={ref(IS, 'cogs', j)}*{ref('Assumptions', 'dpo', j)}/365")
ssb.line("ocl", "Other current liabilities", NUM, hist=lambda j: f"={ref(BS, 'ocl', j)}", fc=lambda j: f"={ref(IS, 'rev', j)}*{ref('Assumptions', 'ocl_pct', j)}")
nwc = lambda j: f"={ref(SCH, 'rec', j)}+{ref(SCH, 'inv', j)}+{ref(SCH, 'oca', j)}-{ref(SCH, 'ap', j)}-{ref(SCH, 'ocl', j)}"  # noqa: E731
ssb.line("nwc", "Net working capital (non-cash)", NUM, bold=True, top=True, hist=nwc, fc=nwc)
ssb.line("dnwc", "Increase / (decrease) in NWC", NUM, hist=lambda j: None if j == 0 else f"={ref(SCH, 'nwc', j)}-{ref(SCH, 'nwc', prev(j))}",
         fc=lambda j: f"={ref(SCH, 'nwc', j)}-{ref(SCH, 'nwc', prev(j))}")
ssb.hdr("PP&E roll-forward")
ssb.line("ppe_beg", "Beginning net PP&E", NUM, hist=lambda j: None if j == 0 else f"={ref(BS, 'ppe', prev(j))}",
         fc=lambda j: f"={ref(BS, 'ppe', prev(j))}")
ssb.line("capex", "(+) Capex", NUM, hist=lambda j: f"=-{raw('cash_flow_capital_expenditures', j)}",
         fc=lambda j: f"={ref(IS, 'rev', j)}*{ref('Assumptions', 'capex_pct', j)}")
ssb.line("ppe_lease", "(+) New leases (right-of-use assets)", NUM, fc=lambda j: f"={ref(SCH, 'lease_add', j)}")
ssb.line("da", "(-) D&A", NUM, hist=lambda j: f"={ref(IS, 'da', j)}",
         fc=lambda j: (f"=IF(ISNUMBER({ref('Assumptions', 'da_ovr', j)}),{ref(IS, 'rev', j)}*{ref('Assumptions', 'da_ovr', j)},"
                       f"({ref(SCH, 'ppe_beg', j)}{fa_leased_term()})*{ref('Assumptions', 'dep_rate', j)})"),
         note="Forecast: consensus-implied D&A (EBITDA - EBIT) where given, otherwise depreciation rate x beginning PP&E.")
ssb.line("ppe_oth", "(+/-) Other (leases, disposals, FX; history only)", NUM,
         hist=lambda j: None if j == 0 else f"={ref(BS, 'ppe', j)}-{ref(SCH, 'ppe_beg', j)}-{ref(SCH, 'capex', j)}+{ref(SCH, 'da', j)}")
ssb.line("ppe_end", "Ending net PP&E", NUM, bold=True, top=True, hist=lambda j: f"={ref(BS, 'ppe', j)}",
         fc=lambda j: f"={ref(SCH, 'ppe_beg', j)}+{ref(SCH, 'capex', j)}+{ref(SCH, 'ppe_lease', j)}-{ref(SCH, 'da', j)}")
ssb.line("ppe_rev", "  PP&E / revenue", PCT, hist=lambda j: f"=IFERROR({ref(SCH, 'ppe_end', j)}/{ref(IS, 'rev', j)},0)",
         fc=lambda j: f"=IFERROR({ref(SCH, 'ppe_end', j)}/{ref(IS, 'rev', j)},0)")
ssb.hdr("Debt & interest (beginning balances - no circularity)")
ssb.line("debt_beg", "Beginning debt (excl. leases)", NUM,
         hist=lambda j: None if j == 0 else f"={ref(BS, 'std', prev(j))}+{ref(BS, 'ltd', prev(j))}",
         fc=lambda j: f"={ref(BS, 'std', prev(j))}+{ref(BS, 'ltd', prev(j))}")
ssb.line("debt_iss", "(+) Net issuance / (repayment)", NUM, hist=lambda j: f"={ref(CF, 'debt', j)}",
         fc=lambda j: f"={ref('Assumptions', 'debt_iss', j)}")
ssb.line("debt_oth", "(+/-) Other (FX, amortisation; history only)", NUM,
         hist=lambda j: None if j == 0 else f"={ref(SCH, 'debt_end', j)}-{ref(SCH, 'debt_beg', j)}-{ref(SCH, 'debt_iss', j)}")
ssb.line("debt_end", "Ending debt (excl. leases)", NUM, bold=True, top=True,
         hist=lambda j: f"={ref(BS, 'std', j)}+{ref(BS, 'ltd', j)}", fc=lambda j: f"={ref(SCH, 'debt_beg', j)}+{ref(SCH, 'debt_iss', j)}")
ssb.line("int_exp", "Interest expense", NUM, fc=lambda j: f"={ref(SCH, 'debt_beg', j)}*{ref('Assumptions', 'kd', j)}")
ssb.line("cash_beg", "Beginning cash", NUM, fc=lambda j: f"={ref(BS, 'cash', prev(j))}")
ssb.line("int_inc", "Interest income", NUM, fc=lambda j: f"=MAX(0,{ref(SCH, 'cash_beg', j)})*{ref('Assumptions', 'cash_yield', j)}")
ssb.hdr("Leases (IFRS 16; zero under US GAAP)")
ssb.line("lease_beg", "Beginning lease liability", NUM, hist=lambda j: None if j == 0 else f"={ref(BS, 'lease', prev(j))}",
         fc=lambda j: f"={ref(BS, 'lease', prev(j))}")
ssb.line("lease_add", "(+) New leases", NUM, fc=lambda j: f"={ref(IS, 'rev', j)}*{ref('Assumptions', 'lease_add', j)}")
ssb.line("lease_rep", "(-) Principal repaid", NUM, hist=lambda j: f"=-{ref(CF, 'lease_rep', j)}",
         fc=lambda j: f"={ref(SCH, 'lease_beg', j)}*{ref('Assumptions', 'lease_rate', j)}")
ssb.line("lease_oth", "(+/-) New leases & other (history only)", NUM,
         hist=lambda j: None if j == 0 else f"={ref(SCH, 'lease_end', j)}-{ref(SCH, 'lease_beg', j)}+{ref(SCH, 'lease_rep', j)}")
ssb.line("lease_end", "Ending lease liability", NUM, bold=True, top=True, hist=lambda j: f"={ref(BS, 'lease', j)}",
         fc=lambda j: f"={ref(SCH, 'lease_beg', j)}+{ref(SCH, 'lease_add', j)}-{ref(SCH, 'lease_rep', j)}")
ssb.hdr("Shareholders' equity roll-forward")
ssb.line("eq_beg", "Beginning equity", NUM, hist=lambda j: None if j == 0 else f"={ref(BS, 'se', prev(j))}",
         fc=lambda j: f"={ref(BS, 'se', prev(j))}")
ssb.line("ni", "(+) Net income to common", NUM, hist=lambda j: f"={ref(IS, 'ni', j)}", fc=lambda j: f"={ref(IS, 'ni', j)}")
ssb.line("div", "(-) Dividends", NUM, hist=lambda j: f"=-{ref(CF, 'div', j)}",
         fc=lambda j: (f"=IF({sref('ret_basis')}=2,MAX(0,{ref(CF, 'fcf', j)}),MAX(0,{ref(SCH, 'ni', j)}))"
                       f"*{ref('Assumptions', 'payout', j)}"),
         note="Basis switch on Assumptions: % of net income, or % of free cash flow (CFO - capex; no circularity).")
ssb.line("bb_tgt", "  Buybacks at the assumed rate", NUM,
         fc=lambda j: (f"=IF({sref('ret_basis')}=2,MAX(0,{ref(CF, 'fcf', j)}),MAX(0,{ref(SCH, 'ni', j)}))"
                       f"*{ref('Assumptions', 'bb_pct', j)}"))
ssb.line("bb", "(-) Share buybacks", NUM, hist=lambda j: f"=-{ref(CF, 'bb', j)}",
         fc=lambda j: (f"=MIN({ref(SCH, 'bb_tgt', j)},MAX(0,{ref(SCH, 'cash_beg', j)}+{ref(CF, 'cfo', j)}+{ref(CF, 'cfi', j)}"
                       f"+{ref(CF, 'debt', j)}+{ref(CF, 'lease_rep', j)}-{ref(SCH, 'div', j)}))"),
         note="Forecast: the assumed rate, capped at the cash left after operations, investing, debt, leases and dividends, "
              "so buybacks never drive cash below zero (no circularity: none of those depend on buybacks).")
ssb.line("eq_iss", "(+) Share issuance", NUM, hist=lambda j: f"={ref(CF, 'iss', j)}")
ssb.line("eq_oth", "(+/-) Other (SBC, OCI, FX; history only)", NUM,
         hist=lambda j: (None if j == 0 else f"={ref(SCH, 'eq_end', j)}-{ref(SCH, 'eq_beg', j)}-{ref(SCH, 'ni', j)}"
                         f"+{ref(SCH, 'div', j)}+{ref(SCH, 'bb', j)}-{ref(SCH, 'eq_iss', j)}"))
ssb.line("eq_end", "Ending equity", NUM, bold=True, top=True, hist=lambda j: f"={ref(BS, 'se', j)}",
         fc=lambda j: f"={ref(SCH, 'eq_beg', j)}+{ref(SCH, 'ni', j)}-{ref(SCH, 'div', j)}-{ref(SCH, 'bb', j)}")
ssb.line("mi_beg", "Beginning minority interest", NUM, hist=lambda j: None if j == 0 else f"={ref(BS, 'mi_bs', prev(j))}",
         fc=lambda j: f"={ref(BS, 'mi_bs', prev(j))}")
ssb.line("mi_inc", "(+) Minority share of income", NUM, hist=lambda j: f"=-{ref(IS, 'mi', j)}", fc=lambda j: f"=-{ref(IS, 'mi', j)}")
ssb.line("mi_oth", "(+/-) Other (minority dividends, FX; history only)", NUM,
         hist=lambda j: None if j == 0 else f"={ref(SCH, 'mi_end', j)}-{ref(SCH, 'mi_beg', j)}-{ref(SCH, 'mi_inc', j)}")
ssb.line("mi_end", "Ending minority interest", NUM, bold=True, top=True, hist=lambda j: f"={ref(BS, 'mi_bs', j)}",
         fc=lambda j: f"={ref(SCH, 'mi_beg', j)}+{ref(SCH, 'mi_inc', j)}")
ssb.close()

# ---------------------------------------------------------------- DCF
DCF = "DCF"
dsb = SB(DCF, "Discounted Cash Flow - unlevered FCF (FCFF)", period_header=False)
W, B, MK, D = M["wacc"], M["bridge"], M["market"], M["dcf"]
vdate = dt.date.fromisoformat(M["valuation_date"])
bdate = dt.date.fromisoformat(B["as_of"])
dsb.hdr("Market & bridge inputs")
dsb.scalar("vdate", "Valuation date", vdate, DATE, source="Date the model was built")
dsb.scalar("bdate", "Balance sheet date for equity bridge", bdate, DATE, source=B.get("source", "Distilla latest reported balance sheet"))
dsb.scalar("price", "Share price", MK["price"], PS, unit=MK.get("price_currency", CUR),
           source=f"Distilla stock_price close {MK.get('price_date', '')}")
dsb.scalar("fx", f"FX: {CUR} per 1 {MK.get('price_currency', CUR)}", MK.get("fx_reporting_per_price", 1.0), '0.0000',
           source=MK.get("fx_source", "1.0 = same currency"))
dsb.scalar("shares", "Shares for value per share", B["diluted_shares"], NUM1, unit="m",
           source=B.get("shares_source") or f"Distilla diluted shares, {B['as_of']}")
dsb.scalar("b_cash", "(+) Cash & ST investments", B["cash"], NUM, source=f"Distilla, {B['as_of']}")
dsb.scalar("b_lti", "(+) Long-term investments (non-operating)", B.get("lt_investments", 0), NUM,
           source="Distilla total LT investments; toggle below if these are operating assets")
dsb.scalar("b_lti_on", "Include LT investments? (1 = yes, 0 = no)", M.get("include_lt_investments", 1), '0', key_fill=True)
dsb.scalar("b_debt", "(-) Total debt excl. leases", B.get("st_debt", 0) + B.get("lt_debt", 0), NUM, source=f"Distilla ST + LT debt, {B['as_of']}")
dsb.scalar("b_lease", "(-) Lease liabilities", B.get("leases", 0), NUM, source=B.get("leases_source", "Distilla"))
dsb.scalar("b_lease_on", "Subtract lease liabilities? (1 = yes, 0 = no)", M.get("subtract_leases", 1), '0', key_fill=True,
           source=(f"{M.get('accounting_standard', '?')}: " + ("lease cost is inside EBIT, so not subtracted again"
                   if not M.get("subtract_leases", 1) else "IFRS 16-style - lease liabilities are debt")))
dsb.scalar("b_mi", "(-) Minority interest", B.get("minority_interest", 0), NUM, source=f"Distilla, {B['as_of']}")
dsb.scalar("b_pens_pre", "Pension & retiree-benefit deficit, pre-tax (blank/0 = none found)", B.get("pension_deficit"), NUM,
           source=B.get("pension_source") or "Not in Distilla; annual report pension note - not sourced")
dsb.scalar("b_pens", "(-) Pension & retiree-benefit deficit, after tax", None, NUM,
           source="Deficit x (1 - tax rate): contributions are tax-deductible; a surplus is not added")
dsb.scalar("target", "Consensus target price (memo)", MK.get("target_price") or 0, PS, unit=MK.get("price_currency", CUR),
           source="Distilla stock_price.sell_side_target_price")
if FA_ON:
    dsb.blank()
    dsb.hdr(f"Finance arm valued separately - {FA.get('name', '')}")
    dsb.scalar("fa_on", "Value the finance arm separately? (1 = yes, 0 = consolidated DCF)", 1, '0', key_fill=True,
               source=f"Segment assets {FA.get('assets', 0):,.0f} = {FA.get('asset_share', 0):.0%} of total assets")
    dsb.scalar("fa_profit", "Finance arm pre-tax profit, last year", FA.get("profit", 0), NUM,
               source=f"{FA.get('segment_source', 'Distilla by_segment_financials')} ({FA.get('period') or 'last year'})")
    dsb.scalar("fa_pct", "  as % of consolidated revenue (held in the forecast)", None, PCT2)
    dsb.scalar("fa_rec_share", "Finance receivables share of current receivables", round(FA.get("rec_share", 0), 4), PCT,
               source=FA.get("rec_share_basis"))
    dsb.scalar("fa_debt", "Finance arm debt (out of the bridge)", round(FA.get("debt", 0), 1), NUM, source=FA.get("debt_basis"))
    dsb.scalar("fa_lti", "Long-term finance receivables inside LT investments", round(FA.get("lt_rec", 0), 1), NUM,
               source=FA.get("lt_rec_basis"))
    dsb.scalar("fa_cash", "Finance arm's own cash (out of bridge cash)", round(FA.get("cash", 0), 1), NUM, source=FA.get("cash_basis"))
    dsb.scalar("fa_leased", "Assets leased to customers, held flat in PP&E", round(FA.get("leased", 0), 1), NUM,
               source=FA.get("leased_basis"))
    dsb.scalar("fa_eq", "Finance arm book equity", round(FA.get("equity", 0), 1), NUM, source=FA.get("equity_basis"))
    dsb.scalar("fa_roe", "Finance arm after-tax ROE", None, PCT, source="Pre-tax profit x (1 - tax rate) / book equity")
    dsb.scalar("fa_pb_ovr", "Your P/B for the finance arm (optional; blank = justified P/B)", FA.get("pb_override"), DEC)
    dsb.scalar("fa_pb", "Finance arm P/B used", None, DEC, bold=True,
               source="Justified P/B = (ROE - g) / (cost of equity - g), bounded 0.5-2.5x; 1.0x if ROE or equity is missing")
    dsb.scalar("fa_val", "Finance arm value (added in the bridge)", None, NUM, bold=True)
dsb.blank()
dsb.hdr("WACC build")
dsb.scalar("rf", "Risk-free rate", W["rf"], PCT2, source=W.get("rf_source"), key_fill=True)
dsb.scalar("beta", "Levered beta", W["beta"], DEC, source=W.get("beta_source"), key_fill=True)
dsb.scalar("erp", "Equity risk premium", W["erp"], PCT2, source=W.get("erp_source"), key_fill=True)
dsb.scalar("crp", "Country risk premium", W["crp"], PCT2, source=W.get("crp_source"))
dsb.scalar("ke", "Cost of equity", None, PCT2, bold=True)
dsb.scalar("kd", "Pre-tax cost of debt", W["kd_pretax"], PCT2, source=W.get("kd_pretax_source"))
dsb.scalar("t", "Tax rate", W["tax_rate"], PCT2, source="Historical effective rate (Distilla) unless overridden")
dsb.scalar("kd_at", "After-tax cost of debt", None, PCT2)
dsb.scalar("mcap", "Market capitalisation", None, NUM, unit=f"{CUR} {UNITS}", source="Price x FX x diluted shares (not Distilla market_cap)")
dsb.scalar("dtot", "Debt incl. leases", None, NUM, unit=f"{CUR} {UNITS}")
dsb.scalar("tgt_wd", "Target debt / capital (blank = current market)", W.get("target_debt_weight"), PCT, key_fill=False)
dsb.scalar("wd", "Weight of debt", None, PCT)
dsb.scalar("we", "Weight of equity", None, PCT)
dsb.scalar("wacc", "WACC", None, PCT2, bold=True, key_fill=True)
dsb.blank()
dsb.hdr("Terminal value assumptions")
dsb.scalar("g", "Terminal growth rate", D["terminal_growth"], PCT2, key_fill=True, source="Perpetuity growth (nominal); keep at or below the risk-free rate")
dsb.scalar("mult", "Exit EV/EBITDA - your multiple (optional; blank = not used)", D.get("exit_multiple"), MULT,
           source=D.get("exit_multiple_basis") or "Enter only with a basis (e.g. mature peers). The valuation uses perpetuity growth.")
dsb.scalar("tv_method", "TV method (1 = perpetuity growth; 2 = your exit multiple, if entered)", D.get("tv_method", 1), '0', key_fill=True)
dsb.scalar("mid", "Mid-year convention (1 = on, 0 = off)", D.get("mid_year", 1), '0')
dsb.blank()
dsb.hdr("Free cash flow build")
TS_HDR_ROW = dsb.r
dsb.r += 1
LAST = NH + NF - 1
fye_dates = [dt.date.fromisoformat(p) for p in P]
dsb.line("fye", "Fiscal year end", DATE, hist=lambda j: fye_dates[j] if j == NH - 1 else None, fc=lambda j: fye_dates[j])
dsb.line("start", "Cash flows counted from", DATE, fc=lambda j: f"=MAX({sref('bdate')},{ref(DCF, 'fye', prev(j))})",
         note="Cash generated before the bridge balance sheet date is already in the cash / debt figures.")
dsb.line("frac", "Fraction of year included", DEC,
         fc=lambda j: f"=MAX(0,MIN(1,({ref(DCF, 'fye', j)}-{ref(DCF, 'start', j)})/({ref(DCF, 'fye', j)}-{ref(DCF, 'fye', prev(j))})))")
dsb.line("ebit", "EBIT", NUM, fc=lambda j: f"={ref(IS, 'ebit', j)}")
if FA_ON:
    dsb.line("fa_ebit", "(-) Finance arm pre-tax profit (valued separately)", NUM,
             fc=lambda j: f"=-{ref(IS, 'rev', j)}*{sref('fa_pct')}*{sref('fa_on')}")
    ebit_in = lambda j: f"({ref(DCF, 'ebit', j)}+{ref(DCF, 'fa_ebit', j)})"  # noqa: E731
else:
    ebit_in = lambda j: ref(DCF, 'ebit', j)  # noqa: E731
dsb.line("tax", "(-) Taxes on EBIT", NUM, fc=lambda j: f"=-{ebit_in(j)}*{ref('Assumptions', 'tax', j)}")
dsb.line("nopat", "NOPAT", NUM, bold=True, top=True, fc=lambda j: f"={ebit_in(j)}+{ref(DCF, 'tax', j)}")
dsb.line("da", "(+) D&A", NUM, fc=lambda j: f"={ref(IS, 'da', j)}")
dsb.line("capex", "(-) Capex", NUM, fc=lambda j: f"=-{ref(SCH, 'capex', j)}")
dsb.line("dnwc", "(-) Increase in NWC", NUM, fc=lambda j: f"=-{ref(SCH, 'dnwc', j)}")
dsb.line("leasecap", "(-) New leases (IFRS 16 right-of-use additions)", NUM, fc=lambda j: f"=-{ref(SCH, 'lease_add', j)}",
         note="New leases are economically capex financed by lease debt. Existing leases are covered by subtracting the lease liability in the bridge.")
if FA_ON:
    dsb.line("fa_rec", "(+) Finance receivables growth (funded by finance-arm debt)", NUM,
             fc=lambda j: f"=({ref(BS, 'rec', j)}-{ref(BS, 'rec', prev(j))})*{sref('fa_rec_share')}*{sref('fa_on')}",
             note="Current finance receivables grow with revenue inside working capital; the finance arm funds them with its own debt, "
                  "which is out of the bridge, so their growth comes out of the industrial cash flow.")
fa_rec_term = (lambda j: f"+{ref(DCF, 'fa_rec', j)}") if FA_ON else (lambda j: "")  # noqa: E731
dsb.line("ufcf", "Unlevered free cash flow", NUM, bold=True, top=True,
         fc=lambda j: (f"={ref(DCF, 'nopat', j)}+{ref(DCF, 'da', j)}+{ref(DCF, 'capex', j)}+{ref(DCF, 'dnwc', j)}"
                       f"+{ref(DCF, 'leasecap', j)}{fa_rec_term(j)}"))
dsb.line("ebitda_v", "EBITDA for multiples" + (" (industrial)" if FA_ON else ""), NUM,
         fc=lambda j: f"={ref(IS, 'ebitda', j)}" + (f"+{ref(DCF, 'fa_ebit', j)}" if FA_ON else ""),
         hist=lambda j: (f"={ref(IS, 'ebitda', j)}" + (f"-{sref('fa_profit')}*{sref('fa_on')}" if FA_ON else ""))
         if j == NH - 1 else None)
dsb.line("rest", "  of which: D&A - capex - NWC - new leases", NUM, fc=lambda j: f"={ref(DCF, 'ufcf', j)}-{ref(DCF, 'nopat', j)}")
dsb.line("ufcf_in", "UFCF counted (x fraction)", NUM, fc=lambda j: f"={ref(DCF, 'ufcf', j)}*{ref(DCF, 'frac', j)}")
dsb.line("tt", "Discount period (years from valuation date)", YRS,
         fc=lambda j: (f"=IF({sref('mid')}=1,({ref(DCF, 'start', j)}+({ref(DCF, 'fye', j)}-{ref(DCF, 'start', j)})/2"
                       f"-{sref('vdate')})/365.25,({ref(DCF, 'fye', j)}-{sref('vdate')})/365.25)"))
dsb.line("df", "Discount factor", '0.0000', fc=lambda j: f"=1/(1+{sref('wacc')})^{ref(DCF, 'tt', j)}")
dsb.line("pv", "PV of UFCF", NUM, bold=True, fc=lambda j: f"={ref(DCF, 'ufcf_in', j)}*{ref(DCF, 'df', j)}")
dsb.blank()
dsb.hdr("Valuation output")
dsb.scalar("sum_pv", "Sum of PV of UFCF", None, NUM)
dsb.scalar("tN", "Discount period for terminal value (years)", None, YRS)
dsb.scalar("tv_g", "Terminal value - perpetuity growth", None, NUM)
dsb.scalar("tv_x", "Terminal value - your exit multiple (if entered)", None, NUM)
dsb.scalar("tv", "Terminal value - selected", None, NUM, bold=True)
dsb.scalar("pv_tv", "PV of terminal value", None, NUM)
dsb.scalar("ev", "Enterprise value", None, NUM, bold=True)
dsb.scalar("bridge", "Net bridge adjustments (cash + investments - debt - leases - MI - pensions"
           + (" + finance arm)" if FA_ON else ")"), None, NUM)
dsb.scalar("eqv", "Equity value", None, NUM, bold=True)
dsb.scalar("vps_rep", f"Equity value per share ({CUR})", None, PS)
dsb.scalar("vps", f"Equity value per share ({MK.get('price_currency', CUR)})", None, PS, bold=True, key_fill=True)
dsb.scalar("px", "Current share price", None, PS)
dsb.scalar("upside", "Upside / (downside)", None, PCT, bold=True)
dsb.scalar("tv_pct", "Terminal value as % of EV", None, PCT)
dsb.scalar("vps_g", "Value per share - perpetuity growth method", None, PS)
dsb.scalar("vs_target", "Value vs consensus target price", None, PCT)
dsb.blank()
dsb.hdr("Cross-check: terminal multiple (a check, not a second valuation)")
dsb.scalar("impl_mult", "Terminal EV/EBITDA implied by the perpetuity method", None, MULT)
dsb.scalar("mkt_mult", "Today's market EV / FY1 EBITDA (reference only)", None, MULT,
           source="Today's multiple prices today's growth; applying it to a mature terminal year is inconsistent")
dsb.scalar("impl_g_mkt", "Perpetual growth implied if today's multiple held at the terminal year", None, PCT2)
dsb.scalar("vps_x", "Value per share at your exit multiple (only if entered)", None, PS)
_mh_src = (f"Distilla valuation_multiple {MH.get('type')} since {MH.get('from')}, n = {MH.get('n')}; "
           f"vendor EV basis: {MH.get('basis') or 'not stated - spot-check against a rebuild'}"
           if MH.get("avg") is not None else "Not retrieved")
dsb.scalar("mh_avg", "Own-history NTM EV/EBITDA - average", float(MH["avg"]) if MH.get("avg") is not None else None, MULT, source=_mh_src)
dsb.scalar("mh_min", "Own-history NTM EV/EBITDA - low", float(MH["min"]) if MH.get("min") is not None else None, MULT)
dsb.blank()
dsb.hdr("Discount rate vs broker notes (the user picks; the model's rate is above)")
_bwm, _bkm = W.get("broker_wacc_median"), W.get("broker_ke_median")
dsb.scalar("b_rate", "Brokers' median discount rate" + (" (cost of equity)" if _bkm and not _bwm else " (WACC)"),
           _bwm if _bwm is not None else _bkm, PCT2, source=W.get("broker_rates") or "None stated in the notes read")
dsb.scalar("b_wacc", "  as a WACC at today's weights", None, PCT2)
dsb.scalar("vps_b", "Value per share at the brokers' rate", None, PS, bold=True)
dsb.scalar("mh_max", "Own-history NTM EV/EBITDA - high", float(MH["max"]) if MH.get("max") is not None else None, MULT)
dsb.blank()
dsb.hdr("Reverse DCF: what today's share price implies")
dsb.scalar("ev_req", "Enterprise value implied by the share price", None, NUM)
dsb.scalar("g_req", "Perpetual growth needed to justify the price (other inputs unchanged)", None, PCT2, bold=True,
           source="Holds terminal cash flow fixed (no extra reinvestment for extra growth), so the true requirement is higher")
dsb.scalar("nopv", "PV of NOPAT, forecast + terminal (perpetuity basis)", None, NUM)
dsb.scalar("k_req", "EBIT margin scale needed to justify the price (1.00 = model)", None, DEC, bold=True,
           source="Approximation: scales EBIT in every forecast year; other drivers unchanged")
dsb.scalar("m_req", "Terminal EBIT margin needed to justify the price", None, PCT)
dsb.blank()
dsb.hdr("Terminal-year reinvestment check")
dsb.scalar("reinv", "Net reinvestment (capex + NWC increase + new leases - D&A)", None, NUM)
dsb.scalar("reinv_rate", "Reinvestment rate (net reinvestment / NOPAT)", None, PCT)
dsb.scalar("ronic", "Implied return on new capital (terminal growth / reinvestment rate)", None, PCT,
           source="Growth needs reinvestment. Below WACC: growth destroys value. No net reinvestment: growth is 'free' - check it")


def fill_dcf_formulas():
    """Scalars that depend on the time-series rows (now that all rows are registered)."""
    CURRENT[0] = DCF
    f0, fN = NH, LAST
    uN = ref(DCF, "ufcf", fN)
    fa_debt_term = f"-{sref('fa_debt')}*{sref('fa_on')}" if FA_ON else ""
    fa_lti_term = f"-{sref('fa_lti')}*{sref('fa_on')}" if FA_ON else ""
    fa_val_term = f"+{sref('fa_val')}*{sref('fa_on')}" if FA_ON else ""
    fa_cash_term = f"-{sref('fa_cash')}*{sref('fa_on')}" if FA_ON else ""
    ebitdaN = ref(DCF, "ebitda_v", fN)
    formulas = {
        "ke": f"={sref('rf')}+{sref('beta')}*{sref('erp')}+{sref('crp')}",
        "kd_at": f"={sref('kd')}*(1-{sref('t')})",
        "mcap": f"={sref('price')}*{sref('fx')}*{sref('shares')}",
        "dtot": f"={sref('b_debt')}{fa_debt_term}+{sref('b_lease')}*{sref('b_lease_on')}",
        "wd": f"=IF({sref('tgt_wd')}=\"\",{sref('dtot')}/({sref('dtot')}+{sref('mcap')}),{sref('tgt_wd')})",
        "we": f"=1-{sref('wd')}",
        "wacc": f"={sref('we')}*{sref('ke')}+{sref('wd')}*{sref('kd_at')}",
        "sum_pv": f"=SUM({rng(DCF, 'pv', f0, fN)})",
        "tN": f"=({ref(DCF, 'fye', fN)}-{sref('vdate')})/365.25",
        "tv_g": f"={uN}*(1+{sref('g')})/({sref('wacc')}-{sref('g')})",
        "tv_x": f"=IF(ISNUMBER({sref('mult')}),{ebitdaN}*{sref('mult')},\"n/a\")",
        "tv": f"=IF(AND({sref('tv_method')}=2,ISNUMBER({sref('mult')})),{sref('tv_x')},{sref('tv_g')})",
        "pv_tv": f"={sref('tv')}/(1+{sref('wacc')})^{sref('tN')}",
        "ev": f"={sref('sum_pv')}+{sref('pv_tv')}",
        "bridge": (f"={sref('b_cash')}{fa_cash_term}+({sref('b_lti')}{fa_lti_term})*{sref('b_lti_on')}-({sref('b_debt')}{fa_debt_term})"
                   f"-{sref('b_lease')}*{sref('b_lease_on')}-{sref('b_mi')}-{sref('b_pens')}{fa_val_term}"),
        "b_pens": f"=MAX(0,N({sref('b_pens_pre')}))*(1-{sref('t')})",
        "eqv": f"={sref('ev')}+{sref('bridge')}",
        "vps_rep": f"=IFERROR({sref('eqv')}/{sref('shares')},0)",
        "vps": f"=IFERROR({sref('vps_rep')}/{sref('fx')},0)",
        "px": f"={sref('price')}",
        "upside": f"=IFERROR({sref('vps')}/{sref('px')}-1,0)",
        "tv_pct": f"=IFERROR({sref('pv_tv')}/{sref('ev')},0)",
        "impl_mult": f"=IFERROR({sref('tv_g')}/{ebitdaN},0)",
        "mkt_mult": f"=IFERROR(({sref('mcap')}-{sref('bridge')})/{ref(DCF, 'ebitda_v', f0)},0)",
        "impl_g_mkt": (f"=IFERROR(({ebitdaN}*{sref('mkt_mult')}*{sref('wacc')}-{uN})/({ebitdaN}*{sref('mkt_mult')}+{uN}),0)"),
        "ev_req": f"={sref('mcap')}-{sref('bridge')}",
        # If the price is below the PV of the forecast years alone, no terminal growth rate solves it
        "g_req": (f"=IF({uN}<=0,\"n/a - terminal cash flow is not positive\","
                  f"IF({sref('ev_req')}<={sref('sum_pv')},\"n/a - price is below the value of the forecast years alone\","
                  f"IF((({sref('ev_req')}-{sref('sum_pv')})*(1+{sref('wacc')})^{sref('tN')}*{sref('wacc')}-{uN})"
                  f"/(({sref('ev_req')}-{sref('sum_pv')})*(1+{sref('wacc')})^{sref('tN')}+{uN})>={sref('wacc')},"
                  f"\"n/a - no growth rate below WACC justifies the price\","
                  f"(({sref('ev_req')}-{sref('sum_pv')})*(1+{sref('wacc')})^{sref('tN')}*{sref('wacc')}-{uN})"
                  f"/(({sref('ev_req')}-{sref('sum_pv')})*(1+{sref('wacc')})^{sref('tN')}+{uN}))))"),
        "nopv": (f"=SUMPRODUCT({rng(DCF, 'nopat', f0, fN)}*{rng(DCF, 'frac', f0, fN)}/(1+{sref('wacc')})^{rng(DCF, 'tt', f0, fN)})"
                 f"+{ref(DCF, 'nopat', fN)}*(1+{sref('g')})/({sref('wacc')}-{sref('g')})/(1+{sref('wacc')})^{sref('tN')}"),
        "k_req": (f"=IF({sref('nopv')}<=0,\"n/a - forecast NOPAT is not positive\","
                  f"IF(1+({sref('ev_req')}-({sref('sum_pv')}+{sref('tv_g')}/(1+{sref('wacc')})^{sref('tN')}))/{sref('nopv')}<=0,"
                  f"\"n/a - no positive margin scale justifies the price\","
                  f"1+({sref('ev_req')}-({sref('sum_pv')}+{sref('tv_g')}/(1+{sref('wacc')})^{sref('tN')}))/{sref('nopv')}))"),
        "m_req": f"=IF(ISNUMBER({sref('k_req')}),{sref('k_req')}*{ref(IS, 'ebit_m', fN)},\"n/a\")",
        "reinv": (f"=-{ref(DCF, 'capex', fN)}-{ref(DCF, 'dnwc', fN)}-{ref(DCF, 'leasecap', fN)}-{ref(DCF, 'da', fN)}"),
        "reinv_rate": f"=IFERROR({sref('reinv')}/{ref(DCF, 'nopat', fN)},0)",
        "ronic": (f"=IF({ref(DCF, 'nopat', fN)}<=0,\"n/a - terminal NOPAT is not positive\","
                  f"IF({sref('reinv_rate')}>0,{sref('g')}/{sref('reinv_rate')},\"n/a - no net reinvestment\"))"),
        "vps_g": (f"=IFERROR(({sref('sum_pv')}+{sref('tv_g')}/(1+{sref('wacc')})^{sref('tN')}+{sref('bridge')})"
                  f"/{sref('shares')}/{sref('fx')},0)"),
        "vps_x": (f"=IF(ISNUMBER({sref('mult')}),({sref('sum_pv')}+{ebitdaN}*{sref('mult')}/(1+{sref('wacc')})^{sref('tN')}"
                  f"+{sref('bridge')})/{sref('shares')}/{sref('fx')},\"n/a\")"),
        "vs_target": f"=IFERROR({sref('vps')}/{sref('target')}-1,0)",
        "b_wacc": (f"=IF(ISNUMBER({sref('b_rate')})," + (f"{sref('we')}*{sref('b_rate')}+{sref('wd')}*{sref('kd_at')}"
                   if (W.get('broker_ke_median') and not W.get('broker_wacc_median')) else f"{sref('b_rate')}") + ",\"n/a\")"),
        "vps_b": (f"=IF(AND(ISNUMBER({sref('b_wacc')}),N({sref('b_wacc')})>{sref('g')}),"
                  f"(SUMPRODUCT({rng(DCF, 'ufcf_in', f0, fN)}/(1+{sref('b_wacc')})^{rng(DCF, 'tt', f0, fN)})"
                  f"+{uN}*(1+{sref('g')})/({sref('b_wacc')}-{sref('g')})/(1+{sref('b_wacc')})^{sref('tN')}"
                  f"+{sref('bridge')})/{sref('shares')}/{sref('fx')},\"n/a\")"),
    }
    if FA_ON:
        lastrev = ref(IS, 'rev', NH - 1)
        formulas.update({
            "fa_pct": f"=IFERROR({sref('fa_profit')}/{lastrev},0)",
            "fa_roe": f"=IFERROR({sref('fa_profit')}*(1-{sref('t')})/{sref('fa_eq')},\"n/a\")",
            "fa_pb": (f"=IF(ISNUMBER({sref('fa_pb_ovr')}),{sref('fa_pb_ovr')},IF(AND(ISNUMBER({sref('fa_roe')}),{sref('fa_eq')}>0),"
                      f"MAX(0.5,MIN(2.5,({sref('fa_roe')}-{sref('g')})/({sref('ke')}-{sref('g')}))),1))"),
            "fa_val": f"={sref('fa_eq')}*{sref('fa_pb')}",
        })
    for kind, r, it in dsb.items:
        if kind == "scalar" and it["key"] in formulas:
            it["value"] = formulas[it["key"]]


# ---------------------------------------------------------------- Sensitivity
SEN = "Sensitivity"
sen = SB(SEN, "Sensitivity - value per share", periods=False)
W_OFF = [-0.01, -0.005, 0, 0.005, 0.01]
G_OFF = [-0.01, -0.005, 0, 0.005, 0.01]


def sens_grid(ws, top, title, col_offsets, col_base_key, col_fmt, method):
    CURRENT[0] = SEN
    ws.cell(row=top, column=1, value=title).font = BOLD
    ws.cell(row=top + 1, column=1, value="WACC (rows) vs " + ("terminal growth" if method == "g" else "exit multiple") +
            " (columns); base case highlighted").font = Font(name=FONT, size=9, italic=True)
    hr = top + 2
    for k, off in enumerate(col_offsets):
        c = ws.cell(row=hr, column=3 + k, value=f"={sref(col_base_key)}+({off})")
        c.number_format, c.font, c.fill = col_fmt, GREEN, SEC_FILL
    f0, fN = NH, LAST
    for i, woff in enumerate(W_OFF):
        r = hr + 1 + i
        wc = ws.cell(row=r, column=2, value=f"={sref('wacc')}+({woff})")
        wc.number_format, wc.font, wc.fill = PCT, GREEN, SEC_FILL
        for k in range(len(col_offsets)):
            w = f"$B{r}"
            x = f"{get_column_letter(3 + k)}${hr}"
            pv = f"SUMPRODUCT({rng(DCF, 'ufcf_in', f0, fN)}/(1+{w})^{rng(DCF, 'tt', f0, fN)})"
            if method == "g":
                tv = f"{ref(DCF, 'ufcf', fN)}*(1+{x})/({w}-{x})"
            else:
                tv = f"{ref(IS, 'ebitda', fN)}*{x}"
            fml = (f"=IFERROR(({pv}+{tv}/(1+{w})^{sref('tN')}+{sref('bridge')})/{sref('shares')}/{sref('fx')},0)")
            c = ws.cell(row=r, column=3 + k, value=fml)
            c.number_format, c.font = PS, BLACK
            if i == 2 and k == 2:
                c.fill, c.font = KEY_FILL, BOLD
    for k in range(len(col_offsets)):
        ws.column_dimensions[get_column_letter(3 + k)].width = 13


sen.custom(lambda ws: sens_grid(ws, 5, "Perpetuity growth method", G_OFF, "g", PCT2, "g"))
K_SC = [0.8, 0.9, 1.0, 1.1, 1.2]


def margin_grid(ws, top):
    """EBIT margin scale (rows) x terminal growth (cols) at the model WACC - an operating sensitivity."""
    CURRENT[0] = SEN
    ws.cell(row=top, column=1, value="EBIT margin x terminal growth (at model WACC)").font = BOLD
    ws.cell(row=top + 1, column=1, value="Rows scale EBIT in every forecast year (1.0 = model); columns = terminal growth").font = \
        Font(name=FONT, size=9, italic=True)
    hr = top + 2
    for k, off in enumerate(G_OFF):
        c = ws.cell(row=hr, column=3 + k, value=f"={sref('g')}+({off})")
        c.number_format, c.font, c.fill = PCT2, GREEN, SEC_FILL
    f0, fN, w = NH, LAST, sref('wacc')
    for i, kk in enumerate(K_SC):
        r = hr + 1 + i
        kc = ws.cell(row=r, column=2, value=kk)
        kc.number_format, kc.font, kc.fill = '0.00"x"', BLUE, SEC_FILL
        for k in range(len(G_OFF)):
            x, kr = f"{get_column_letter(3 + k)}${hr}", f"$B{r}"
            pv = (f"SUMPRODUCT(({kr}*{rng(DCF, 'nopat', f0, fN)}+{rng(DCF, 'rest', f0, fN)})*{rng(DCF, 'frac', f0, fN)}"
                  f"/(1+{w})^{rng(DCF, 'tt', f0, fN)})")
            tv = f"({kr}*{ref(DCF, 'nopat', fN)}+{ref(DCF, 'rest', fN)})*(1+{x})/({w}-{x})"
            c = ws.cell(row=r, column=3 + k,
                        value=f"=IFERROR(({pv}+{tv}/(1+{w})^{sref('tN')}+{sref('bridge')})/{sref('shares')}/{sref('fx')},0)")
            c.number_format, c.font = PS, BLACK
            if i == 2 and k == 2:
                c.fill, c.font = KEY_FILL, BOLD


sen.custom(lambda ws: margin_grid(ws, 15))

# ---------------------------------------------------------------- Checks
CHK = "Checks"
TH = M.get("thresholds", {})
chk = SB(CHK, "Model integrity checks")
chk.hdr("Every year")
chk.line("bs", "Balance sheet balances (assets - L&E)", NUM, hist=lambda j: f"={ref(BS, 'check', j)}",
         fc=lambda j: f"={ref(BS, 'check', j)}")
chk.line("cash_neg", "Forecast cash below zero? (1 = yes)", '0', fc=lambda j: f"=IF({ref(BS, 'cash', j)}<0,1,0)")
chk.line("cf_tie", "Cash flow ending cash - balance sheet cash", NUM, fc=lambda j: f"=ROUND({ref(CF, 'cash_end', j)}-{ref(BS, 'cash', j)},3)")
chk.hdr("History ties to Distilla (model sum - Distilla reported; all should be 0)")
chk.line("t_cfo", "Cash from operations", NUM, hist=lambda j: f"=ROUND({ref(CF, 'cfo', j)}-{raw('cash_flow_net_operating_cash_flow', j)},3)")
chk.line("t_cfi", "Cash from investing", NUM, hist=lambda j: f"=ROUND({ref(CF, 'cfi', j)}-{raw('cash_flow_net_investing_cash_flow', j)},3)")
chk.line("t_cff", "Cash from financing", NUM, hist=lambda j: f"=ROUND({ref(CF, 'cff', j)}-{raw('cash_flow_net_financing_cash_flow', j)},3)")
chk.line("t_net", "Net change in cash", NUM, hist=lambda j: f"=ROUND({ref(CF, 'net', j)}-{raw('cash_flow_net_change_in_cash', j)},3)")
chk.line("t_bs", "Balance sheet unreconciled lines (sum of absolute values)", NUM,
         hist=lambda j: "=ROUND(" + "+".join(f"ABS({ref(BS, k, j)})" for k in ("u_ca", "u_lta", "u_cl", "u_ltl", "u_eq")) + ",3)")
chk.line("t_ni", "Net income to common", NUM, hist=lambda j: f"=ROUND({ref(IS, 'ni', j)}-{raw('income_statement_net_income', j)},3)")
chk.hdr("Forecast shape")
chk.line("ppe_rev", "PP&E / revenue", PCT, hist=lambda j: f"={ref(SCH, 'ppe_rev', j)}", fc=lambda j: f"={ref(SCH, 'ppe_rev', j)}")
chk.line("cash_rev", "Cash / revenue", PCT, hist=lambda j: f"=IFERROR({ref(BS, 'cash', j)}/{ref(IS, 'rev', j)},0)",
         fc=lambda j: f"=IFERROR({ref(BS, 'cash', j)}/{ref(IS, 'rev', j)},0)")
chk.line("opex_neg", "Operating expenses negative? (1 = yes)", '0', fc=lambda j: f"=IF({ref(IS, 'opex', j)}<0,1,0)")
chk.line("bb_cap", "Buybacks capped by available cash? (1 = yes)", '0',
         fc=lambda j: f"=IF({ref(SCH, 'bb', j)}<{ref(SCH, 'bb_tgt', j)}-0.5,1,0)")
chk.blank()
chk.hdr("Summary")
chk.scalar("c_bs", "Max absolute balance sheet difference", None, NUM)
chk.scalar("c_cash", "Years with negative cash", None, '0')
chk.scalar("c_hist", "History ties to Distilla (CF subtotals, BS, net income)", None, '@')
chk.scalar("c_wg", "WACC exceeds terminal growth", None, '@')
chk.scalar("c_tv", f"Terminal value share of EV below {TH.get('tv_share_warn', 0.85):.0%}", None, '@')
chk.scalar("c_capex", "Terminal-year capex + new leases at least D&A", None, '@')
chk.scalar("c_ppe", "Terminal PP&E / revenue within range of last actual", None, '@')
chk.scalar("c_cashb", "No excess cash build-up", None, '@')
chk.scalar("c_opex", "Operating expenses never negative", None, '@')
chk.scalar("c_eqv", "Equity value positive", None, '@')
chk.scalar("c_bbcap", "Years with buybacks capped by available cash (information)", None, '0')
if FA_ON:
    chk.scalar("c_fa", "Finance arm inputs within the consolidated figures", None, '@')
chk.scalar("c_reinv", "Terminal growth consistent with reinvestment (RONIC >= WACC; no large 'free' growth)", None, '@')
chk.scalar("c_all", "OVERALL", None, '@', bold=True, key_fill=True)


def fill_checks():
    CURRENT[0] = CHK
    lo, hi, cb = TH.get("ppe_drift_low", 0.5), TH.get("ppe_drift_high", 2.0), TH.get("cash_build_multiple", 3.0)
    hist_rows = ("t_cfo", "t_cfi", "t_cff", "t_net", "t_bs", "t_ni")
    maxabs = "+".join(f"MAX(ABS(MAX({rng(CHK, k, 0, NH - 1)})),ABS(MIN({rng(CHK, k, 0, NH - 1)})))" for k in hist_rows)
    ppeN, ppe0 = ref(CHK, 'ppe_rev', LAST), ref(CHK, 'ppe_rev', NH - 1)
    cashN, cash0 = ref(CHK, 'cash_rev', LAST), ref(CHK, 'cash_rev', NH - 1)
    f = {
        "c_bs": f"=MAX(ABS(MAX({rng(CHK, 'bs', 0, LAST)})),ABS(MIN({rng(CHK, 'bs', 0, LAST)})))",
        "c_cash": f"=SUM({rng(CHK, 'cash_neg', NH, LAST)})",
        "c_hist": f"=IF(({maxabs})<MAX(1,0.0001*{ref(IS, 'rev', NH - 1)}),\"OK\",\"REVIEW\")",
        "c_wg": f"=IF({sref('wacc')}>{sref('g')},\"OK\",\"FAIL\")",
        "c_tv": f"=IF({sref('tv_pct')}<{TH.get('tv_share_warn', 0.85)},\"OK\",\"REVIEW\")",
        "c_capex": f"=IF({ref(SCH, 'capex', LAST)}+{ref(SCH, 'ppe_lease', LAST)}>={ref(SCH, 'da', LAST)},\"OK\",\"REVIEW\")",
        "c_ppe": f"=IF(AND({ppeN}>={lo}*{ppe0},{ppeN}<={hi}*{ppe0}),\"OK\",\"REVIEW\")",
        "c_cashb": f"=IF(AND({cashN}>{cb}*{cash0},{cashN}>0.5),\"REVIEW\",\"OK\")",
        "c_opex": f"=IF(SUM({rng(CHK, 'opex_neg', NH, LAST)})=0,\"OK\",\"FAIL\")",
        "c_eqv": f"=IF({sref('eqv')}>0,\"OK\",\"REVIEW\")",
        "c_bbcap": f"=SUM({rng(CHK, 'bb_cap', NH, LAST)})",
        "c_reinv": (f"=IF(OR({sref('g')}<=0,{ref(DCF, 'nopat', LAST)}<=0),\"OK\",IF({sref('reinv_rate')}<-0.1,\"REVIEW\","
                    f"IF({sref('reinv_rate')}<=0,\"OK\",IF({sref('g')}/{sref('reinv_rate')}<{sref('wacc')},\"REVIEW\",\"OK\"))))"),
        # hard failures = the model is wrong; warnings = the model is right but an assumption needs a look
        "c_all": (f"=IF(AND({sref('c_bs')}<1,{sref('c_wg')}=\"OK\",{sref('c_opex')}=\"OK\"),"
                  f"IF(AND({sref('c_cash')}=0,{sref('c_tv')}=\"OK\",{sref('c_capex')}=\"OK\",{sref('c_hist')}=\"OK\","
                  f"{sref('c_ppe')}=\"OK\",{sref('c_cashb')}=\"OK\",{sref('c_eqv')}=\"OK\",{sref('c_reinv')}=\"OK\"),\"ALL CHECKS PASS\",\"PASS WITH WARNINGS\"),"
                  f"\"FAIL - REVIEW\")"),
    }
    if FA_ON:
        f["c_fa"] = (f"=IF(AND({sref('fa_debt')}<={sref('b_debt')},{sref('fa_lti')}<={sref('b_lti')},{sref('fa_eq')}>0,"
                     f"{sref('fa_cash')}<={sref('b_cash')},{sref('fa_leased')}<={ref(BS, 'ppe', NH - 1)},"
                     f"{sref('fa_rec_share')}<=1),\"OK\",\"REVIEW\")")
        f["c_all"] = f["c_all"].replace(f"{sref('c_reinv')}=\"OK\")", f"{sref('c_reinv')}=\"OK\",{sref('c_fa')}=\"OK\")")
    for kind, r, it in chk.items:
        if kind == "scalar" and it["key"] in f:
            it["value"] = f[it["key"]]


# ---------------------------------------------------------------- Summary
SUM_ = "Summary"
sm = SB(SUM_, "Valuation Summary", periods=False)
sm.text("Company", M["company"].get("name"), bold=True)
sm.text("Ticker", M["company"].get("ticker"))
sm.text("Reporting currency / units", f"{CUR} {UNITS} (per-share values in {MK.get('price_currency', CUR)})")
sm.text("Data source", "Distilla MCP (actuals, consensus, prices); WACC inputs per DCF tab sources")
sm.blank()
sm.hdr("Valuation")
sm.scalar("s_scen", "Active scenario", None, '@')
sm.scalar("s_vps", "Equity value per share", None, PS, bold=True, key_fill=True, unit=MK.get("price_currency", CUR))
sm.scalar("s_px", "Current share price", None, PS, unit=MK.get("price_currency", CUR))
sm.scalar("s_up", "Upside / (downside)", None, PCT, bold=True)
sm.scalar("s_tgt", "Consensus target price", None, PS, unit=MK.get("price_currency", CUR))
sm.scalar("s_ev", "Enterprise value", None, NUM, unit=f"{CUR} {UNITS}")
sm.scalar("s_eq", "Equity value", None, NUM, unit=f"{CUR} {UNITS}")
sm.scalar("s_wacc", "WACC", None, PCT2)
sm.scalar("s_g", "Terminal growth", None, PCT2)
sm.scalar("s_tvp", "Terminal value % of EV", None, PCT)
sm.scalar("s_vg", "Value / share - perpetuity method", None, PS)
sm.scalar("s_greq", "Share price implies perpetual growth of", None, PCT2, bold=True)
sm.scalar("s_mreq", "... or a terminal EBIT margin of", None, PCT, bold=True)
sm.scalar("s_im", "Terminal EV/EBITDA implied by the valuation", None, MULT)
sm.scalar("s_mkt", "Today's market EV/EBITDA (reference)", None, MULT)
if FA_ON:
    sm.scalar("s_fa", f"Finance arm value in the bridge ({FA.get('name', '')})", None, NUM, unit=f"{CUR} {UNITS}")
    sm.scalar("s_fapb", "  at P/B (justified unless overridden; 1.0x = default)", None, DEC)
if W.get("broker_rates"):
    sm.scalar("s_brate", "Brokers' median discount rate (as WACC)", None, PCT2)
    sm.scalar("s_vpsb", "Value per share at the brokers' rate", None, PS, bold=True, unit=MK.get("price_currency", CUR))
if MH.get("avg") is not None:
    sm.scalar("s_mh", "Own 3-year NTM EV/EBITDA average (reference)", None, MULT)
sm.scalar("s_chk", "Model checks", None, '@', bold=True)
sm.blank()
sm.hdr("Forecast snapshot")
SNAP_TOP = sm.r
sm.r += 6
BASIS = M.get("basis") or {}
if BASIS:
    sm.hdr("Key assumptions - where each number comes from")
    b_top = sm.r
    brows = [(drv, sc, BASIS[drv][sc]) for drv in ("revenue_growth", "ebit_margin") for sc in ("base", "bull", "bear")
             if sc in BASIS.get(drv, {})]
    sm.r += len(brows) + 1

    def _basis(ws):
        for c, h in zip((1, 2, 4, 5), ("Driver", "Scenario", "Basis", "Type")):
            x = ws.cell(row=b_top, column=c, value=h)
            x.font, x.fill = HDR_FONT, HDR_FILL
        for i, (drv, sc, b) in enumerate(brows):
            r = b_top + 1 + i
            vals = {1: drv.replace("_", " "), 2: sc, 4: b.get("text", ""), 5: b.get("type", "")}
            for c, v in vals.items():
                x = ws.cell(row=r, column=c, value=v)
                x.font = Font(name=FONT, size=9, color=("C00000" if "no evidence" in str(b.get("text", "")) and c == 4 else "000000"))
                x.alignment = Alignment(wrap_text=True, vertical="top")
    sm.custom(_basis)
    sm.blank()
EV_ = M.get("evidence") or []
if EV_:
    sm.hdr("Assumption evidence (qualitative sources, paraphrased)")
    ev_top = sm.r
    sm.r += len(EV_) + 1

    def _evidence(ws):
        heads = ("Assumption", "Stance", "Finding", "Source", "Date", "Effect on the numbers")
        cols = (1, 2, 4, 5, 6, 7)  # finding goes in the wide column D
        for c, h in zip(cols, heads):
            x = ws.cell(row=ev_top, column=c, value=h)
            x.font, x.fill = HDR_FONT, HDR_FILL
        for i, e in enumerate(EV_):
            r = ev_top + 1 + i
            for c, k in zip(cols, ("assumption", "stance", "finding", "source", "date", "effect")):
                x = ws.cell(row=r, column=c, value=e.get(k, ""))
                x.font = Font(name=FONT, size=9)
                x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["E"].width = 34
        ws.column_dimensions["F"].width = 12
        ws.column_dimensions["G"].width = 40
    sm.custom(_evidence)
    sm.blank()
sm.hdr("Model notes - deliberate simplifications (small effect on the DCF)")
for note in (
        "Stock-based compensation is treated as a real cost: not added back in forecast cash flow.",
        "Long-term investments are held flat and earn no forecast income; their value is added in the equity bridge.",
        "Equity-method (affiliate) income is not forecast; the investment value sits in long-term investments.",
        "Intangibles, deferred taxes, provisions, tax payable and other long-term items are held flat.",
        "All forecast D&A runs through PP&E; book tax = cash tax; no dividends to minority holders.",
        "Forecast acquisitions, disposals, investment purchases, share issuance and FX are zero (history shows them).",
        "Share count forecast is used for EPS only; valuation uses the latest diluted share count."):
    sm.text("  " + note)
sm.blank()
sm.hdr("Flags & data notes")
for fl in M.get("flags", []) or ["None"]:
    sm.text("  " + fl)


def fill_summary():
    CURRENT[0] = SUM_
    f = {"s_scen": f"='Assumptions'!D{sel_row}",
         "s_vps": f"={sref('vps')}", "s_px": f"={sref('px')}", "s_up": f"={sref('upside')}",
         "s_tgt": f"={sref('target')}", "s_ev": f"={sref('ev')}", "s_eq": f"={sref('eqv')}",
         "s_wacc": f"={sref('wacc')}", "s_g": f"={sref('g')}", "s_tvp": f"={sref('tv_pct')}",
         "s_vg": f"={sref('vps_g')}", "s_greq": f"={sref('g_req')}", "s_mreq": f"={sref('m_req')}", "s_im": f"={sref('impl_mult')}",
         "s_mkt": f"={sref('mkt_mult')}",
         "s_chk": f"={sref('c_all')}"}
    if FA_ON:
        f.update({"s_fa": f"={sref('fa_val')}*{sref('fa_on')}", "s_fapb": f"={sref('fa_pb')}"})
    if MH.get("avg") is not None:
        f["s_mh"] = f"={sref('mh_avg')}"
    if W.get("broker_rates"):
        f.update({"s_brate": f"={sref('b_wacc')}", "s_vpsb": f"={sref('vps_b')}"})
    for kind, r, it in sm.items:
        if kind == "scalar" and it["key"] in f:
            it["value"] = f[it["key"]]

    def snap(ws):
        CURRENT[0] = SUM_
        rows = [("Revenue", IS, "rev", NUM), ("Revenue growth", IS, "growth", PCT), ("EBIT margin", IS, "ebit_m", PCT),
                ("Net income", IS, "ni", NUM), ("Unlevered FCF", DCF, "ufcf", NUM)]
        for k, j in enumerate(range(NH - 1, LAST + 1)):
            c = ws.cell(row=SNAP_TOP, column=3 + k, value=fy_label(P[j], j))
            c.font, c.fill = HDR_FONT, HDR_FILL
            ws.column_dimensions[get_column_letter(3 + k)].width = 13
        for i, (lab, sh, key, fmt) in enumerate(rows):
            r = SNAP_TOP + 1 + i
            ws.cell(row=r, column=1, value=lab).font = BLACK
            for k, j in enumerate(range(NH - 1, LAST + 1)):
                if sh == DCF and is_hist(j):
                    continue
                c = ws.cell(row=r, column=3 + k, value=f"={ref(sh, key, j)}")
                c.number_format, c.font = fmt, GREEN
    sm.custom(snap)


# ---------------------------------------------------------------- write
fill_dcf_formulas()
fill_checks()
fill_summary()

wb = Workbook()
wb.remove(wb.active)
for sb in ((sm, asb, drv, model, dsb, sen, chk, rawsb) if DRV else (sm, asb, model, dsb, sen, chk, rawsb)):
    ws = write_sheet(wb, sb)
    if sb is dsb:
        for j in range(len(P)):
            c = ws.cell(row=TS_HDR_ROW, column=FIRST_COL + j, value=fy_label(P[j], j))
            c.font, c.fill = HDR_FONT, HDR_FILL
    if sb is sen:
        CURRENT[0] = SEN
        ws.cell(row=24, column=1, value="Values per share in " + MK.get("price_currency", CUR) + "; current price:")
        c = ws.cell(row=24, column=3, value=f"={sref('px')}")
        c.number_format, c.font = PS, GREEN

wb["Summary"].sheet_properties.tabColor = "1F3864"
wb["Assumptions"].sheet_properties.tabColor = "FFC000"
wb["Checks"].sheet_properties.tabColor = "70AD47"
if DRV:
    wb["Drivers"].sheet_properties.tabColor = "FFC000"
wb.save(OUT)
print(f"Wrote {OUT}: {NH} historical + {NF} forecast years, sheets: {', '.join(wb.sheetnames)}")
